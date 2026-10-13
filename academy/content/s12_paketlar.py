from dsl import SECTION, TOPIC, P, Q, F, PR, L

SECTION(12, "Mashhur paketlar (amaliy)", "API, SMS, web scraping, avtomatlashtirish, PDF, Excel va NumPy", "🚀")

BIZ = '''{"businesses": [
 {"name": "Afsona", "rating": 4.7, "review_count": 812, "price": "$$", "distance": 1250.5, "categories": [{"title": "Uzbek"}], "location": {"city": "Toshkent"}},
 {"name": "Caravan", "rating": 4.5, "review_count": 1530, "price": "$$$", "distance": 3400.0, "categories": [{"title": "Uzbek"}, {"title": "Cafe"}], "location": {"city": "Toshkent"}},
 {"name": "Beshqozon", "rating": 4.8, "review_count": 2950, "price": "$", "distance": 8900.2, "categories": [{"title": "Plov"}], "location": {"city": "Toshkent"}},
 {"name": "Kofe Uyi", "rating": 4.2, "review_count": 95, "price": "$", "distance": 450.0, "categories": [{"title": "Cafe"}], "location": {"city": "Toshkent"}}
], "total": 4}'''

# ---------------------------------------------------------------- 12.1
TOPIC("api-nima", "API nima? HTTP asoslari", "Soʻrov, javob, metodlar va status kodlari", 20, r'''
**API** (Application Programming Interface) — dasturlarning bir-biri bilan "gaplashish" usuli. Ob-havo ilovasi ob-havo serveridan, Telegram bot Telegram serveridan, toʻlov tizimi bankdan maʼlumotni API orqali oladi.

Internetdagi API'lar koʻpincha **HTTP** orqali ishlaydi: siz **soʻrov** (request) yuborasiz, server **javob** (response) qaytaradi — odatda JSON formatida.

```text
SOʻROV  →  GET https://api.example.com/v1/weather?city=Toshkent
           Authorization: Bearer <kalit>

JAVOB   ←  200 OK
           Content-Type: application/json
           {"city": "Toshkent", "temp": 24, "sky": "ochiq"}
```

## HTTP metodlari
| Metod | Maʼnosi | Misol |
|---|---|---|
| `GET` | maʼlumot olish | mahsulotlar roʻyxati |
| `POST` | yangi yaratish | buyurtma berish |
| `PUT` / `PATCH` | oʻzgartirish | profilni yangilash |
| `DELETE` | oʻchirish | izohni oʻchirish |

## Status kodlari
```py
from http import HTTPStatus

for kod in [200, 201, 400, 401, 404, 429, 500]:
    s = HTTPStatus(kod)
    print(kod, s.phrase)
```

| Diapazon | Maʼnosi |
|---|---|
| 2xx | ✅ muvaffaqiyat |
| 3xx | ↪ yoʻnaltirish |
| 4xx | ❌ **siz** xato qildingiz (notoʻgʻri soʻrov, kalit yoʻq, topilmadi) |
| 5xx | 💥 **server**da xato |

## requests — eng mashhur HTTP kutubxona
```py norun
import requests                      # pip install requests

r = requests.get("https://api.github.com/repos/python/cpython", timeout=10)
print(r.status_code)                 # 200
data = r.json()                      # JSON → lugʻat
print(data["stargazers_count"])

r = requests.post("https://httpbin.org/post", json={"ism": "Ali"}, timeout=10)
print(r.json()["json"])              # {'ism': 'Ali'}
```

> `timeout=` ni doim bering — aks holda server javob bermasa, dasturingiz abadiy kutib qolishi mumkin.

## API'ni oʻzimiz simulyatsiya qilamiz
Haqiqiy server oʻrniga oddiy funksiya — lekin mantiq xuddi shunday:

```py
import json

MAHSULOTLAR = {1: {"nom": "Non", "narx": 4000}, 2: {"nom": "Sut", "narx": 12000}}

def api(metod, yol, tana=None):
    qismlar = yol.strip("/").split("/")
    if qismlar[0] != "mahsulotlar":
        return 404, {"xato": "topilmadi"}
    if metod == "GET" and len(qismlar) == 1:
        return 200, list(MAHSULOTLAR.values())
    if metod == "GET":
        m = MAHSULOTLAR.get(int(qismlar[1]))
        return (200, m) if m else (404, {"xato": "mahsulot yoʻq"})
    if metod == "POST":
        yangi_id = max(MAHSULOTLAR) + 1
        MAHSULOTLAR[yangi_id] = tana
        return 201, {"id": yangi_id}
    return 405, {"xato": "metod ruxsat etilmagan"}

for soʻrov in [("GET", "/mahsulotlar"), ("GET", "/mahsulotlar/2"), ("GET", "/mahsulotlar/9"),
               ("POST", "/mahsulotlar", {"nom": "Choy", "narx": 25000}), ("DELETE", "/mahsulotlar/1")]:
    kod, javob = api(*soʻrov)
    print(kod, json.dumps(javob, ensure_ascii=False))
```

## Esda tuting
- API = soʻrov (metod + URL + sarlavhalar + tana) → javob (status + JSON).
- 4xx — mijoz xatosi, 5xx — server xatosi.
''', [
    P('from http import HTTPStatus\nprint(HTTPStatus(404).phrase)', "404", "Error", "Not Here"),
    P('from http import HTTPStatus\nprint(HTTPStatus.CREATED.value)', "CREATED", "200", "Created"),
    P('kod = 503\nprint("server" if kod >= 500 else "mijoz" if kod >= 400 else "ok")', "mijoz", "ok", "503"),
    P('import json\nprint(json.loads(\'{"temp": 24}\')["temp"] + 1)', "241", "Xato", "temp"),
    P('from urllib.parse import urlencode\nprint("https://api.uz/v1/ob-havo?" + urlencode({"shahar": "Xiva", "kun": 3}))', "https://api.uz/v1/ob-havo?shahar=Xiva,kun=3", "https://api.uz/v1/ob-havo?{'shahar': 'Xiva'}", "Xato"),
    P('from http import HTTPStatus\nprint(HTTPStatus(429).phrase)', "Rate Limit", "Error", "Too Fast", lv=2),
    P('print(200 <= 204 < 300, 200 <= 302 < 300)', "True True", "False False", "False True"),
    Q("Yangi buyurtma yaratish uchun odatda qaysi HTTP metodi ishlatiladi?", "POST", "GET", "DELETE", "HEAD"),
    Q("`401 Unauthorized` nimani bildiradi?", "Autentifikatsiya (kalit yoki token) yoʻq yoki notoʻgʻri", "Server ishlamayapti", "Sahifa topilmadi", "Hammasi muvaffaqiyatli"),
    F("Status turi", "`status_turi(kod)` — 2xx → `\"muvaffaqiyat\"`, 3xx → `\"yoʻnaltirish\"`, 4xx → `\"mijoz xatosi\"`, 5xx → `\"server xatosi\"`, aks holda `\"nomaʼlum\"`.",
      "def status_turi(kod):\n    return {2: 'muvaffaqiyat', 3: 'yoʻnaltirish', 4: 'mijoz xatosi', 5: 'server xatosi'}.get(kod // 100, 'nomaʼlum')",
      "status_turi(200)", "status_turi(404)", "status_turi(503)", "status_turi(301)", "status_turi(99)", lv=1),
    F("Status matni", "`matn(kod)` — `http.HTTPStatus` dan `\"404 Not Found\"` koʻrinishidagi satr; nomaʼlum kod — `\"KOD Unknown\"`.",
      "from http import HTTPStatus\n\n\ndef matn(kod):\n    try:\n        return f'{kod} {HTTPStatus(kod).phrase}'\n    except ValueError:\n        return f'{kod} Unknown'",
      "matn(404)", "matn(201)", "matn(799)"),
    F("URL yasash", "`url(asos, yol, **parametrlar)` — `asos` + `yol` + (parametr boʻlsa) `?` + `urlencode` (parametrlar alifbo boʻyicha).",
      "from urllib.parse import urlencode\n\n\ndef url(asos, yol, **parametrlar):\n    u = asos.rstrip('/') + '/' + yol.lstrip('/')\n    return u + ('?' + urlencode(sorted(parametrlar.items())) if parametrlar else '')",
      'url("https://api.uz/", "/v1/ob-havo", shahar="Xiva", kun=3)', 'url("https://api.uz", "mahsulotlar")', 'url("http://localhost:8000", "/q", s="a b")'),
    F("Mini API", "`api(metod, yol)` — `/talabalar` (GET → hammasi roʻyxati), `/talabalar/<id>` (GET → bitta yoki 404). `TALABALAR = {1: \"Ali\", 2: \"Madina\"}`. Javob — `(status, maʼlumot)`; boshqa metod — `(405, None)`; boshqa yoʻl — `(404, None)`.",
      "TALABALAR = {1: 'Ali', 2: 'Madina'}\n\n\ndef api(metod, yol):\n    q = yol.strip('/').split('/')\n    if q[0] != 'talabalar' or len(q) > 2:\n        return (404, None)\n    if metod != 'GET':\n        return (405, None)\n    if len(q) == 1:\n        return (200, list(TALABALAR.values()))\n    t = TALABALAR.get(int(q[1])) if q[1].isdigit() else None\n    return (200, t) if t else (404, None)",
      'api("GET", "/talabalar")', 'api("GET", "/talabalar/2")', 'api("GET", "/talabalar/9")', 'api("POST", "/talabalar")', 'api("GET", "/kurslar")', lv=2),
    F("Javobni tekshirish", "`ishlov(status, tana)` — 200/201 boʻlsa JSON matnni lugʻatga aylantirib qaytarsin; 4xx/5xx boʻlsa `RuntimeError(\"HTTP KOD: xabar\")` chiqarsin (xabar — JSON dagi `\"xato\"` maydoni, boʻlmasa `\"nomaʼlum\"`).",
      "import json\n\n\ndef ishlov(status, tana):\n    d = json.loads(tana) if tana else {}\n    if status >= 400:\n        raise RuntimeError(f\"HTTP {status}: {d.get('xato', 'nomaʼlum')}\")\n    return d",
      "ishlov(200, '{\"ism\": \"Ali\"}')", "ishlov(404, '{\"xato\": \"topilmadi\"}')", "ishlov(500, '')", "ishlov(201, '{}')", lv=2),
    F("Qayta urinish", "`urin(javoblar, max_urinish)` — server javoblari (status kodlar) ketma-ketligi. 5xx yoki 429 boʻlsa qayta urinsin; 2xx — `(\"ok\", urinishlar)`, 4xx (429 dan tashqari) — `(\"xato\", urinishlar)`, urinishlar tugasa — `(\"taslim\", max_urinish)`.",
      "def urin(javoblar, max_urinish):\n    for i, kod in enumerate(javoblar[:max_urinish], 1):\n        if 200 <= kod < 300:\n            return ('ok', i)\n        if 400 <= kod < 500 and kod != 429:\n            return ('xato', i)\n    return ('taslim', max_urinish)",
      "urin([503, 503, 200], 5)", "urin([500, 404], 5)", "urin([429, 429, 429], 3)", "urin([200], 1)", lv=3),
    F("JSON tana", "`buyurtma_tana(mahsulotlar)` — `[(nom, soni)]` dan POST uchun JSON satr: `{\"items\": [{\"nom\": ..., \"soni\": ...}], \"jami_dona\": N}` (`ensure_ascii=False`).",
      "import json\n\n\ndef buyurtma_tana(mahsulotlar):\n    items = [{'nom': n, 'soni': s} for n, s in mahsulotlar]\n    return json.dumps({'items': items, 'jami_dona': sum(s for _, s in mahsulotlar)}, ensure_ascii=False)",
      'buyurtma_tana([("non", 2), ("sut", 1)])', "buyurtma_tana([])", 'buyurtma_tana([("oʻrik", 5)])'),
    L("docs", "requests-quickstart"),
    L("docs", "py-http-status"),
    L("exercism", "rest-api"),
    L("docs", "py-urllib-request", title="urllib.request — tashqi paketsiz HTTP"),
])

# ---------------------------------------------------------------- 12.2
TOPIC("yelp-api", "Haqiqiy API: biznes qidiruvi (Yelp)", "Kalit bilan soʻrov, parametrlar va javobni qayta ishlash", 25, r'''
Kafe, restoran yoki doʻkonlarni qidiradigan xizmatlar (Yelp, Google Places, 2GIS) oʻz API'larini taqdim etadi. Ishlash tartibi deyarli bir xil:

1. Saytda roʻyxatdan oʻtib **API kalit** olasiz;
2. Har bir soʻrovga kalitni **sarlavha**da qoʻshasiz;
3. Qidiruv shartlarini **parametr** sifatida berasiz;
4. Kelgan **JSON** ni qayta ishlaysiz.

## Yelp Fusion API bilan qidiruv
```py norun
import os
import requests

KALIT = os.environ["YELP_API_KEY"]            # kalit koddan tashqarida!
url = "https://api.yelp.com/v3/businesses/search"
sarlavhalar = {"Authorization": f"Bearer {KALIT}"}
parametrlar = {"term": "plov", "location": "Tashkent", "limit": 10, "sort_by": "rating"}

javob = requests.get(url, headers=sarlavhalar, params=parametrlar, timeout=10)
javob.raise_for_status()                       # 4xx/5xx boʻlsa xato chiqaradi
for b in javob.json()["businesses"]:
    print(b["name"], b["rating"])
```

## Javob tuzilishi
API javobi odatda shunday koʻrinadi (qisqartirilgan):

```py
import json

javob = """{"businesses": [
 {"name": "Afsona", "rating": 4.7, "review_count": 812, "price": "$$", "distance": 1250.5},
 {"name": "Caravan", "rating": 4.5, "review_count": 1530, "price": "$$$", "distance": 3400.0},
 {"name": "Beshqozon", "rating": 4.8, "review_count": 2950, "price": "$", "distance": 8900.2}
], "total": 3}"""

data = json.loads(javob)
print("Topildi:", data["total"])
for b in sorted(data["businesses"], key=lambda b: -b["rating"]):
    print(f"{b['name']:<10} ⭐{b['rating']}  {b['review_count']:>5} sharh  {b['distance'] / 1000:.1f} km  {b['price']}")
```

## Filtrlash va reyting
```py
import json

data = json.loads('{"businesses": [{"name": "A", "rating": 4.7, "review_count": 812}, {"name": "B", "rating": 4.9, "review_count": 12}, {"name": "C", "rating": 4.5, "review_count": 1530}]}')

# Kam sharhli 4.9 dan koʻp sharhli 4.7 ishonchliroq: oddiy "ogʻirlikli" reyting
def ball(b):
    return round(b["rating"] * min(1, b["review_count"] / 100), 2)

for b in sorted(data["businesses"], key=ball, reverse=True):
    print(b["name"], ball(b))
```

## Sahifalash (pagination)
API'lar natijalarni boʻlib-boʻlib beradi: `limit` (nechta) va `offset` (qayerdan).

```py
jami, limit = 47, 20
for offset in range(0, jami, limit):
    print(f"offset={offset}, limit={limit}")
```

!> Har bir API'ning **limitlari** bor (masalan, kuniga 500 soʻrov). `429 Too Many Requests` kelsa — kutib, keyin qayta urining.

## Esda tuting
- Kalit — `Authorization: Bearer ...` sarlavhasida, parametrlar — `params=` da.
- `raise_for_status()` — xatoni darhol aniqlash; `timeout=` — doim.
''', [
    P('h = {"Authorization": "Bearer " + "abc123"}\nprint(h["Authorization"].split()[0])', "abc123", "Authorization", "Bearer abc123"),
    P('import json\nd = json.loads(\'{"businesses": [{"name": "A"}, {"name": "B"}], "total": 2}\')\nprint([b["name"] for b in d["businesses"]])', "['A']", "AB", "2"),
    P('print(f"{8900.2 / 1000:.1f} km")', "8900.2 km", "8.9002 km", "9 km"),
    P('print(len(range(0, 47, 20)))', "2", "47", "20"),
    P('print(list(range(0, 47, 20)))', "[20, 40]", "[0, 20, 40, 47]", "[0, 47]"),
    P('b = [{"n": "A", "r": 4.5}, {"n": "B", "r": 4.8}]\nprint(max(b, key=lambda x: x["r"])["n"])', "A", "4.8", "Xato"),
    P('narx = ["$", "$$$", "$$"]\nprint(sorted(narx, key=len))', "['$', '$$$', '$$']", "['$$$', '$$', '$']", "Xato"),
    Q("API kaliti soʻrovning qaysi qismida yuboriladi (Yelp kabi)?", "Authorization sarlavhasida (Bearer token)", "URL ning oxirida ochiq holda", "Javob tanasida", "Fayl nomida"),
    Q("`raise_for_status()` nima qiladi?", "Status 4xx yoki 5xx boʻlsa xato chiqaradi", "Javobni JSON ga aylantiradi", "Soʻrovni qayta yuboradi", "Kalitni tekshiradi"),
    F("Eng yaxshi joylar", "`top(matn, n)` — API javobidan (`businesses`) eng yuqori reytingli n ta biznes nomini qaytarsin; reyting teng boʻlsa sharhlar soni koʻpi oldin.",
      "import json\n\n\ndef top(matn, n):\n    b = json.loads(matn)['businesses']\n    return [x['name'] for x in sorted(b, key=lambda x: (-x['rating'], -x['review_count']))[:n]]",
      "top(__import__('biz').JAVOB, 2)", "top(__import__('biz').JAVOB, 10)", "top('{\"businesses\": []}', 3)", files={"biz.py": "JAVOB = '''" + BIZ + "'''\n"}, lv=1),
    F("Kategoriya boʻyicha", "`kategoriya(matn, nom)` — `categories[].title` ichida `nom` boʻlgan bizneslar nomlari (masofa boʻyicha yaqinidan).",
      "import json\n\n\ndef kategoriya(matn, nom):\n    b = json.loads(matn)['businesses']\n    return [x['name'] for x in sorted(b, key=lambda x: x['distance']) if any(c['title'] == nom for c in x['categories'])]",
      "kategoriya(__import__('biz').JAVOB, 'Cafe')", "kategoriya(__import__('biz').JAVOB, 'Uzbek')", "kategoriya(__import__('biz').JAVOB, 'Sushi')", files={"biz.py": "JAVOB = '''" + BIZ + "'''\n"}),
    F("Hamyonbop va yaqin", "`tanla(matn, max_narx, max_km)` — narx belgisi (`$`) soni `max_narx` dan oshmagan va masofasi `max_km` km dan kam bizneslar nomlari (reyting boʻyicha kamayish).",
      "import json\n\n\ndef tanla(matn, max_narx, max_km):\n    b = json.loads(matn)['businesses']\n    ok = [x for x in b if len(x['price']) <= max_narx and x['distance'] / 1000 < max_km]\n    return [x['name'] for x in sorted(ok, key=lambda x: -x['rating'])]",
      "tanla(__import__('biz').JAVOB, 2, 5)", "tanla(__import__('biz').JAVOB, 1, 100)", "tanla(__import__('biz').JAVOB, 3, 0.1)", files={"biz.py": "JAVOB = '''" + BIZ + "'''\n"}, lv=2),
    F("Soʻrov parametrlari", "`parametrlar(term, joy, limit=20, sahifa=1)` — Yelp uslubidagi lugʻat: `term`, `location`, `limit` (1..50 oraligʻiga siqilsin), `offset` = `(sahifa-1)*limit`.",
      "def parametrlar(term, joy, limit=20, sahifa=1):\n    limit = max(1, min(50, limit))\n    return {'term': term, 'location': joy, 'limit': limit, 'offset': (sahifa - 1) * limit}",
      'parametrlar("plov", "Tashkent")', 'parametrlar("cafe", "Samarkand", 100, 3)', 'parametrlar("x", "y", 0)'),
    F("Sahifalar", "`sahifalar(jami, limit)` — barcha natijalarni olish uchun kerak boʻlgan `(offset, limit)` juftliklari roʻyxati (oxirgisida limit qolgan miqdorga teng).",
      "def sahifalar(jami, limit):\n    return [(o, min(limit, jami - o)) for o in range(0, jami, limit)]",
      "sahifalar(47, 20)", "sahifalar(0, 20)", "sahifalar(40, 20)", lv=2),
    F("Ishonchli reyting", "`ishonchli(rating, sharhlar, m=100, ortacha=4.0)` — Bayes oʻrtachasi: `(sharhlar*rating + m*ortacha) / (sharhlar + m)`, 2 xona.",
      "def ishonchli(rating, sharhlar, m=100, ortacha=4.0):\n    return round((sharhlar * rating + m * ortacha) / (sharhlar + m), 2)",
      "ishonchli(4.9, 12)", "ishonchli(4.7, 812)", "ishonchli(5.0, 0)", lv=3),
    F("Xulosa jadvali", "`jadval(matn)` — har bir biznes uchun `\"NOM ⭐R (N)\"` qatorlari, reyting boʻyicha kamayish, `\\n` bilan birlashtirilgan satr.",
      "import json\n\n\ndef jadval(matn):\n    b = sorted(json.loads(matn)['businesses'], key=lambda x: -x['rating'])\n    return '\\n'.join(f\"{x['name']} ⭐{x['rating']} ({x['review_count']})\" for x in b)",
      "jadval(__import__('biz').JAVOB)", "jadval('{\"businesses\": []}')", "jadval('{\"businesses\": [{\"name\": \"X\", \"rating\": 3.0, \"review_count\": 1}]}')", files={"biz.py": "JAVOB = '''" + BIZ + "'''\n"}),
    L("docs", "requests-quickstart"),
    L("exercism", "rest-api"),
    L("docs", "py-json"),
    L("lc", "design-a-food-rating-system"),
])

# ---------------------------------------------------------------- 12.3
TOPIC("api-kalit", "API kalitlarini yashirish", "Muhit oʻzgaruvchilari, .env fayllar va .gitignore", 15, r'''
API kalit — parol bilan teng. U kod ichida yozilib, GitHub'ga tushsa, bir necha **daqiqada** botlar topib, hisobingizdan foydalanishi mumkin.

!> Hech qachon kalitni koddagi oʻzgaruvchiga yozmang: `KALIT = "sk-12345..."`. Bir marta Git tarixiga tushsa, uni oʻchirish ham yordam bermaydi — kalitni **bekor qilib**, yangisini olish kerak.

## 1-usul: muhit oʻzgaruvchilari
```bash
# Windows (PowerShell)
$env:YELP_API_KEY = "abc123"
# macOS / Linux
export YELP_API_KEY="abc123"
```

```py
import os

os.environ["YELP_API_KEY"] = "abc123def456"       # (bu yerda terminal oʻrniga)
kalit = os.environ.get("YELP_API_KEY")
if not kalit:
    raise SystemExit("YELP_API_KEY oʻrnatilmagan")
print("Kalit yuklandi:", kalit[:3] + "…" + kalit[-3:])
```

## 2-usul: .env fayl
Loyiha papkasida `.env` fayl yaratiladi va **Git'ga qoʻshilmaydi**:

```text
# .env
YELP_API_KEY=abc123def456
TWILIO_TOKEN=xyz789
DEBUG=true
```

```py norun
from dotenv import load_dotenv      # pip install python-dotenv
import os

load_dotenv()                        # .env ni os.environ ga yuklaydi
kalit = os.getenv("YELP_API_KEY")
```

`python-dotenv` aslida juda oddiy ish qiladi — uni oʻzimiz ham yoza olamiz:

```py
import os
from pathlib import Path

Path(".env").write_text('# sozlamalar\nYELP_API_KEY=abc123\nSALOM="Assalomu alaykum"\n\nDEBUG=true\n', encoding="utf-8")

def env_yukla(fayl=".env"):
    for qator in Path(fayl).read_text(encoding="utf-8").splitlines():
        qator = qator.strip()
        if not qator or qator.startswith("#") or "=" not in qator:
            continue
        kalit, qiymat = qator.split("=", 1)
        os.environ.setdefault(kalit.strip(), qiymat.strip().strip('"'))

env_yukla()
print(os.environ["YELP_API_KEY"], "|", os.environ["SALOM"], "|", os.environ["DEBUG"])
```

## .gitignore va .env.example
```text
# .gitignore
.env
```

```text
# .env.example  ← Git'ga qoʻshiladi, qiymatlarsiz namuna
YELP_API_KEY=
TWILIO_TOKEN=
DEBUG=false
```

> GitHub Actions'da kalitlar **Secrets**da saqlanadi (Mirae Academy'ning Telegram bot tokeni ham aynan shunday — kodda yoʻq!).

## Esda tuting
- Kalit — muhit oʻzgaruvchisida yoki `.env` da; `.env` — `.gitignore` da.
- Loglarda kalitni toʻliq chiqarmang — maskalang: `abc…456`.
''', [
    P('import os\nos.environ["K"] = "x"\nprint(os.environ.get("K"), os.environ.get("YOQ_K_123"))', "x ''", "K None", "⚠ KeyError"),
    P('import os\nprint(os.environ["MIRAE_HECH_YOQ_789"])', "None", "''", "MIRAE_HECH_YOQ_789"),
    P('q = "KEY=abc=def"\nprint(q.split("=", 1))', "['KEY', 'abc', 'def']", "('KEY', 'abc=def')", "Xato"),
    P('k = "sk_live_1234567890"\nprint(k[:3] + "…" + k[-3:])', "sk_l…7890", "sk_live…7890", "Xato"),
    P('print(\'"Salom dunyo"\'.strip(\'"\'))', '"Salom dunyo"', "Salom", "Xato"),
    P('import os\nos.environ["A"] = "1"\nos.environ.setdefault("A", "2")\nprint(os.environ["A"])', "2", "12", "Xato"),
    P('print(bool(""), bool("false"))', "False False", "True True", "Xato", lv=2,
      why="Boʻsh boʻlmagan har qanday satr (hatto `\"false\"`) — `True`. Shuning uchun `.env` dagi `DEBUG=false` ni alohida tekshirish kerak."),
    Q("Kalit tasodifan GitHub'ga tushib qolsa nima qilish kerak?", "Kalitni darhol bekor qilib, yangisini olish", "Faylni oʻchirish yetarli", "Repository nomini oʻzgartirish", "Hech narsa — hech kim koʻrmaydi"),
    Q("`.env.example` fayli nima uchun kerak?", "Qaysi oʻzgaruvchilar kerakligini qiymatlarsiz koʻrsatish uchun (Git'ga qoʻshiladi)", "Haqiqiy kalitlarni saqlash uchun", "Python versiyasini koʻrsatish uchun", "Testlar uchun"),
    F("Maskalash", "`maska(kalit)` — 8 belgidan uzun boʻlsa birinchi 4 va oxirgi 4 belgini qoldirib, oʻrtasini `…` bilan almashtirsin; aks holda faqat `\"****\"`.",
      "def maska(kalit):\n    return kalit[:4] + '…' + kalit[-4:] if len(kalit) > 8 else '****'",
      'maska("sk_live_1234567890")', 'maska("abc")', 'maska("123456789")', lv=1),
    F(".env oʻqish", "`env_oqi(matn)` — `.env` matnidan lugʻat: izoh va boʻsh qatorlar tashlansin, `KEY=VALUE` (birinchi `=` boʻyicha), qiymat atrofidagi qoʻshtirnoqlar (`\"` yoki `'`) olib tashlansin.",
      "def env_oqi(matn):\n    d = {}\n    for q in matn.splitlines():\n        q = q.strip()\n        if not q or q.startswith('#') or '=' not in q:\n            continue\n        k, v = q.split('=', 1)\n        v = v.strip()\n        if len(v) >= 2 and v[0] == v[-1] and v[0] in '\"\\'':\n            v = v[1:-1]\n        d[k.strip()] = v\n    return d",
      'env_oqi("# izoh\\nKEY=abc\\nSALOM=\\"Assalomu alaykum\\"\\n\\nURL=https://a.uz/?x=1\\n")', 'env_oqi("")', "env_oqi(\"A='1'\\nB=2\")", lv=2),
    F("Mantiqiy sozlama", "`bool_env(qiymat)` — `\"1\"`, `\"true\"`, `\"yes\"`, `\"on\"` (katta-kichik farqsiz) → `True`; `\"0\"`, `\"false\"`, `\"no\"`, `\"off\"`, `\"\"` → `False`; boshqasi — `ValueError`.",
      "def bool_env(qiymat):\n    q = qiymat.strip().lower()\n    if q in ('1', 'true', 'yes', 'on'):\n        return True\n    if q in ('0', 'false', 'no', 'off', ''):\n        return False\n    raise ValueError(qiymat)",
      'bool_env("TRUE")', 'bool_env("false")', 'bool_env("")', 'bool_env("balki")'),
    F("Yetishmayotganlar", "`yetishmaydi(kerak, muhit)` — `kerak` roʻyxatidagi oʻzgaruvchilardan `muhit` lugʻatida yoʻq yoki boʻsh boʻlganlarini qaytarsin.",
      "def yetishmaydi(kerak, muhit):\n    return [k for k in kerak if not muhit.get(k)]",
      'yetishmaydi(["YELP_API_KEY", "DEBUG"], {"YELP_API_KEY": "abc", "DEBUG": ""})', "yetishmaydi([], {})", 'yetishmaydi(["A"], {})', lv=1),
    F("Kalit sizib chiqqanmi", "`sizgan_qatorlar(kod)` — kod matnidagi qator raqamlari (1 dan): `KEY`, `TOKEN`, `SECRET` yoki `PASSWORD` soʻzi (katta-kichik farqsiz) bor va unga uzunligi 8+ boʻlgan satr literal berilgan qatorlar (`os.environ` / `getenv` bor qatorlar hisobga olinmasin).",
      "import re\n\n\ndef sizgan_qatorlar(kod):\n    natija = []\n    for i, q in enumerate(kod.splitlines(), 1):\n        if re.search(r'(key|token|secret|password)', q, re.I) and re.search(r'=\\s*[\\'\"][^\\'\"]{8,}[\\'\"]', q) and 'environ' not in q and 'getenv' not in q:\n            natija.append(i)\n    return natija",
      'sizgan_qatorlar("import os\\nAPI_KEY = \\"sk_live_abcdefgh\\"\\ntoken = os.getenv(\\"TOKEN\\")\\nPASSWORD = \\"123\\"\\n")', 'sizgan_qatorlar("")', 'sizgan_qatorlar("secret = \'qwertyuiop\'")', lv=3),
    F("Namuna fayl", "`namuna(env_matn)` — `.env` dan `.env.example` yasasin: har bir `KEY=qiymat` → `KEY=`, izoh va boʻsh qatorlar oʻz holicha qolsin.",
      "def namuna(env_matn):\n    q = []\n    for x in env_matn.splitlines():\n        s = x.strip()\n        if s and not s.startswith('#') and '=' in s:\n            q.append(s.split('=', 1)[0] + '=')\n        else:\n            q.append(x)\n    return '\\n'.join(q)",
      'namuna("# Yelp\\nYELP_API_KEY=abc123\\n\\nDEBUG=true")', 'namuna("")', 'namuna("A=1=2")'),
    F("Sozlamalar ustuvorligi", "`sozlama(nom, muhit, env_fayl, standart=None)` — qiymatni avval `muhit` lugʻatidan (haqiqiy muhit oʻzgaruvchilari), topilmasa `env_fayl` lugʻatidan, u yerda ham boʻlmasa `standart` dan olsin.",
      "def sozlama(nom, muhit, env_fayl, standart=None):\n    if nom in muhit:\n        return muhit[nom]\n    if nom in env_fayl:\n        return env_fayl[nom]\n    return standart",
      'sozlama("DEBUG", {"DEBUG": "0"}, {"DEBUG": "1"})', 'sozlama("DEBUG", {}, {"DEBUG": "1"})', 'sozlama("PORT", {}, {}, "8000")', lv=1),
    L("docs", "dotenv"),
    L("docs", "py-os", title="os.environ — rasmiy hujjat"),
    L("docs", "dj-settings", title="Django settings — maxfiy sozlamalar"),
    L("docs", "dj-checklist", title="Deployment checklist: SECRET_KEY"),
])

# ---------------------------------------------------------------- 12.4
TOPIC("sms", "SMS yuborish", "Twilio uslubidagi API, raqam formati va SMS segmentlari", 20, r'''
Tasdiqlash kodi, buyurtma holati, eslatma — SMS hamon eng ishonchli aloqa kanallaridan biri. Python'dan SMS yuborish uchun **SMS-shlyuz** API'si ishlatiladi: xalqaro **Twilio**, Oʻzbekistonda — Eskiz, Play Mobile kabi xizmatlar.

## Twilio bilan
```py norun
import os
from twilio.rest import Client          # pip install twilio

client = Client(os.environ["TWILIO_SID"], os.environ["TWILIO_TOKEN"])
xabar = client.messages.create(
    body="Mirae Academy: tasdiqlash kodingiz 482913",
    from_="+15017122661",                # Twilio bergan raqam
    to="+998901234567",
)
print(xabar.sid, xabar.status)
```

Aslida bu oddiy HTTP POST soʻrovi:

```py norun
import requests, os

r = requests.post(
    f"https://api.twilio.com/2010-04-01/Accounts/{os.environ['TWILIO_SID']}/Messages.json",
    data={"From": "+15017122661", "To": "+998901234567", "Body": "Salom!"},
    auth=(os.environ["TWILIO_SID"], os.environ["TWILIO_TOKEN"]),
    timeout=10,
)
```

## Telefon raqam formati: E.164
Xalqaro standart — `+` va mamlakat kodi bilan, boʻsh joysiz: `+998901234567`.

```py
import re

def e164(raqam):
    r = re.sub(r"\D", "", raqam)              # faqat raqamlar
    if len(r) == 9:                           # 90 123 45 67
        r = "998" + r
    return "+" + r

for x in ["90 123 45 67", "+998 (91) 765-43-21", "998935554433"]:
    print(f"{x:<22} → {e164(x)}")
```

## SMS uzunligi va segmentlar
!> Lotin harflarida bitta SMS — 160 belgi. Lekin matnda `oʻ`, `gʻ`, `ʼ` yoki kirill harflari boʻlsa, telefon **Unicode (UCS-2)** rejimiga oʻtadi — bitta SMS atigi **70** belgi! Uzun xabar boʻlaklarga boʻlinadi va har bir boʻlak alohida pul turadi.

| Kodlash | 1 SMS | Uzun xabar (har boʻlak) |
|---|---|---|
| GSM-7 (lotin) | 160 | 153 |
| UCS-2 (unicode) | 70 | 67 |

```py
import math

GSM = set("@£$¥èéùìòÇ\nØø\rÅåΔ_ΦΓΛΩΠΨΣΘΞÆæßÉ !\"#¤%&'()*+,-./0123456789:;<=>?¡ABCDEFGHIJKLMNOPQRSTUVWXYZÄÖÑÜ§¿abcdefghijklmnopqrstuvwxyzäöñüà")

def segmentlar(matn):
    if set(matn) <= GSM:
        return 1 if len(matn) <= 160 else math.ceil(len(matn) / 153)
    return 1 if len(matn) <= 70 else math.ceil(len(matn) / 67)

print(segmentlar("Kodingiz: 482913"))
print(segmentlar("Hurmatli mijoz, buyurtmangiz yo'lga chiqdi. Kuryer 30 daqiqada yetib boradi."))
print(segmentlar("Hurmatli mijoz, buyurtmangiz yoʻlga chiqdi. Kuryer 30 daqiqada yetib boradi."))
```

> Shu sababli koʻp kompaniyalar SMS matnida `oʻ` oʻrniga oddiy `o'` ishlatadi — narx ikki baravar arzon chiqadi.

## Esda tuting
- Raqamni E.164 ga keltiring: `+998XXXXXXXXX`.
- Unicode belgi bitta boʻlsa ham SMS hajmi 160 dan 70 ga tushadi.
''', [
    P('import re\nprint(re.sub(r"\\D", "", "+998 (90) 123-45-67"))', "+998901234567", "998 90 123 45 67", "Xato"),
    P('print(len("998901234567"))', "13", "9", "11"),
    P('import math\nprint(math.ceil(200 / 153))', "1", "1.3", "3"),
    P('print(set("salom") <= set("abcdefghijklmnopqrstuvwxyz"))', "False", "salom", "Xato"),
    P('print("oʻ" == "o\'", len("oʻ"), len("o\'"))', "True 2 2", "False 1 1", "True 1 1"),
    P('import unicodedata\nprint(unicodedata.name("ʻ"))', "APOSTROPHE", "OKINA", "Xato", lv=2),
    P('t = "+998901234567"\nprint(t.startswith("+998") and len(t) == 13 and t[1:].isdigit())', "False", "+998", "Xato"),
    Q("SMS matnida bitta `oʻ` harfi boʻlsa nima oʻzgaradi?", "Kodlash Unicode'ga oʻtadi va bitta SMS 160 emas, 70 belgi boʻladi", "Hech narsa", "SMS yuborilmaydi", "Matn avtomatik tarjima qilinadi"),
    Q("E.164 formatidagi toʻgʻri raqam qaysi?", "+998901234567", "8 (90) 123-45-67", "90 123 45 67", "998-90-123-45-67"),
    F("E.164", "`e164(raqam)` — barcha belgilarni olib tashlab, 9 xonali boʻlsa `998` qoʻshsin, natija `+998` + 9 raqam boʻlmasa `ValueError`.",
      "import re\n\n\ndef e164(raqam):\n    r = re.sub(r'\\D', '', raqam)\n    if len(r) == 9:\n        r = '998' + r\n    if not (len(r) == 12 and r.startswith('998')):\n        raise ValueError(raqam)\n    return '+' + r",
      'e164("90 123 45 67")', 'e164("+998 (91) 765-43-21")', 'e164("998935554433")', 'e164("12345")', lv=1),
    F("Operator", "`operator(raqam)` — E.164 raqamning kodi (`+998` dan keyingi 2 raqam) boʻyicha: `90`, `91` → `\"Beeline\"`; `93`, `94` → `\"Ucell\"`; `97`, `88` → `\"Mobiuz\"`; `95`, `99` → `\"Uzmobile\"`; `33` → `\"Humans\"`; boshqasi → `\"nomaʼlum\"`.",
      "def operator(raqam):\n    k = raqam[4:6]\n    return {'90': 'Beeline', '91': 'Beeline', '93': 'Ucell', '94': 'Ucell', '97': 'Mobiuz', '88': 'Mobiuz', '95': 'Uzmobile', '99': 'Uzmobile', '33': 'Humans'}.get(k, 'nomaʼlum')",
      'operator("+998901234567")', 'operator("+998931112233")', 'operator("+998771234567")'),
    F("Segmentlar", "`segment(matn)` — matnda faqat ASCII belgilar boʻlsa (soddalashtirilgan GSM): ≤160 → 1, aks holda `ceil(len/153)`; ASCII dan tashqari belgi boʻlsa: ≤70 → 1, aks holda `ceil(len/67)`.",
      "import math\n\n\ndef segment(matn):\n    if matn.isascii():\n        return 1 if len(matn) <= 160 else math.ceil(len(matn) / 153)\n    return 1 if len(matn) <= 70 else math.ceil(len(matn) / 67)",
      'segment("Kodingiz: 482913")', 'segment("a" * 161)', 'segment("oʻ" * 36)', 'segment("x" * 160)', lv=2),
    F("Narx hisobi", "`narx(xabarlar, bir_sms=60)` — har bir xabar segmentlari (`isascii` asosida, yuqoridagi qoida) yigʻindisini bir SMS narxiga koʻpaytirsin.",
      "import math\n\n\ndef narx(xabarlar, bir_sms=60):\n    def seg(m):\n        if m.isascii():\n            return 1 if len(m) <= 160 else math.ceil(len(m) / 153)\n        return 1 if len(m) <= 70 else math.ceil(len(m) / 67)\n    return sum(seg(m) for m in xabarlar) * bir_sms",
      'narx(["Salom", "Kod: 1234"])', 'narx(["oʻ" * 50], 100)', "narx([])", lv=2),
    F("Lotinlashtirish", "`arzonlashtir(matn)` — SMS arzon boʻlishi uchun: `ʻ` va `ʼ` → `'`; natija ASCII boʻlmasa ham qaytarilsin.",
      "def arzonlashtir(matn):\n    return matn.replace('ʻ', \"'\").replace('ʼ', \"'\")",
      'arzonlashtir("Oʻzbekiston, maʼlumot")', 'arzonlashtir("oddiy")', 'arzonlashtir("gʻoʻza").isascii()', lv=1),
    F("Tasdiqlash xabari", "`kod_xabari(kod, daqiqa=5)` — `\"Mirae Academy: tasdiqlash kodingiz KOD. U D daqiqa amal qiladi. Kodni hech kimga bermang.\"`; kod 6 xonali raqam boʻlmasa `ValueError`.",
      "def kod_xabari(kod, daqiqa=5):\n    if not (len(kod) == 6 and kod.isdigit()):\n        raise ValueError(kod)\n    return f'Mirae Academy: tasdiqlash kodingiz {kod}. U {daqiqa} daqiqa amal qiladi. Kodni hech kimga bermang.'",
      'kod_xabari("482913")', 'kod_xabari("000001", 10)', 'kod_xabari("12ab56")'),
    F("Ommaviy yuborish", "`yubor(raqamlar, matn)` — har bir raqamni E.164 ga keltirsin (notoʻgʻrilarini tashlab), takrorlarni olib tashlasin va `{\"yuborildi\": [...], \"xato\": N}` qaytarsin (raqamlar kelish tartibida).",
      "import re\n\n\ndef yubor(raqamlar, matn):\n    ok, xato = [], 0\n    for x in raqamlar:\n        r = re.sub(r'\\D', '', x)\n        if len(r) == 9:\n            r = '998' + r\n        if len(r) == 12 and r.startswith('998'):\n            if '+' + r not in ok:\n                ok.append('+' + r)\n        else:\n            xato += 1\n    return {'yuborildi': ok, 'xato': xato}",
      'yubor(["90 123 45 67", "+998901234567", "123", "+998 93 555 44 33"], "Salom")', 'yubor([], "x")', 'yubor(["1"], "x")', lv=3),
    L("docs", "requests-quickstart", title="Requests: POST va autentifikatsiya"),
    L("exercism", "phone-number", note="Telefon raqamini tozalash va tekshirish — aynan shu mavzu."),
    L("docs", "dotenv"),
    L("lc", "unique-email-addresses"),
])

# ---------------------------------------------------------------- 12.5
TOPIC("web-scraping", "Web scraping", "Sahifalardan maʼlumot yigʻish: requests, BeautifulSoup, html.parser", 25, r'''
**Web scraping** — veb-sahifadagi maʼlumotni dastur yordamida avtomatik yigʻish: narxlar monitoringi, yangiliklar sarlavhalari, vakansiyalar roʻyxati...

## Odatiy toʻplam: requests + BeautifulSoup
```py norun
import requests
from bs4 import BeautifulSoup          # pip install beautifulsoup4

html = requests.get("https://example.com/yangiliklar", timeout=10,
                    headers={"User-Agent": "MiraeBot/1.0 (talaba loyihasi)"}).text
soup = BeautifulSoup(html, "html.parser")

for maqola in soup.select("article.yangilik"):
    sarlavha = maqola.select_one("h2").get_text(strip=True)
    havola = maqola.select_one("a")["href"]
    print(sarlavha, "→", havola)
```

## Kutubxonasiz: html.parser
Python'ning oʻzida ham HTML tahlilchi bor. Uning ishlashini tushunish — BeautifulSoup'ni ham yaxshiroq tushunish demak:

```py
from html.parser import HTMLParser

HTML = """
<html><head><title>Mirae yangiliklari</title></head>
<body>
  <article class="yangilik"><h2>Python 3.14 chiqdi</h2><a href="/py314">Batafsil</a></article>
  <article class="yangilik"><h2>Yangi kurs: Django</h2><a href="/django">Batafsil</a></article>
  <a href="https://mirae.uz">Bosh sahifa</a>
</body></html>
"""

class Yigʻuvchi(HTMLParser):
    def __init__(self):
        super().__init__()
        self.havolalar = []
        self.sarlavhalar = []
        self._h2 = False

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.havolalar.append(dict(attrs).get("href"))
        if tag == "h2":
            self._h2 = True

    def handle_endtag(self, tag):
        if tag == "h2":
            self._h2 = False

    def handle_data(self, data):
        if self._h2:
            self.sarlavhalar.append(data.strip())

y = Yigʻuvchi()
y.feed(HTML)
print(y.sarlavhalar)
print(y.havolalar)
```

## Nisbiy havolalarni toʻliq qilish
```py
from urllib.parse import urljoin

asos = "https://mirae.uz/yangiliklar/"
for h in ["/py314", "django", "https://python.org"]:
    print(urljoin(asos, h))
```

## Odob qoidalari va robots.txt
Sayt egasining qoidalarini hurmat qiling: `robots.txt` da qaysi sahifalarni yigʻish mumkin emasligi yozilgan.

```py
from urllib.robotparser import RobotFileParser

rp = RobotFileParser()
rp.parse("""
User-agent: *
Disallow: /admin/
Disallow: /profil/
Crawl-delay: 2
""".splitlines())
for yol in ["/yangiliklar", "/admin/sozlama", "/profil/ali"]:
    print(yol, rp.can_fetch("MiraeBot", "https://mirae.uz" + yol))
print("Kutish:", rp.crawl_delay("MiraeBot"), "soniya")
```

!> Scraping qilishda: `robots.txt` va foydalanish shartlarini oʻqing, soʻrovlar orasida kuting (`time.sleep`), shaxsiy maʼlumotlarni yigʻmang. Rasmiy API boʻlsa — doim API'ni afzal koʻring.

## Esda tuting
- `requests` sahifani yuklaydi, `BeautifulSoup`/`html.parser` uni tahlil qiladi.
- `urljoin` — nisbiy havolalar, `robots.txt` — ruxsatlar.
''', [
    P('from urllib.parse import urljoin\nprint(urljoin("https://a.uz/x/y", "z"))', "https://a.uz/z", "https://a.uz/x/y/z", "z"),
    P('from urllib.parse import urljoin\nprint(urljoin("https://a.uz/x/y", "/z"))', "https://a.uz/x/z", "https://a.uz/x/y/z", "/z"),
    P('from html.parser import HTMLParser\nclass P(HTMLParser):\n    n = 0\n    def handle_starttag(self, t, a):\n        P.n += t == "p"\np = P()\np.feed("<p>a</p><div><p>b</p></div>")\nprint(P.n)', "1", "3", "0"),
    P('import html\nprint(html.unescape("5 &lt; 7 &amp;&amp; 7 &gt; 5"))', "5 &lt; 7 &amp;&amp; 7 &gt; 5", "5 < 7 & 7 > 5", "Xato"),
    P('import html\nprint(html.escape("<b>Salom</b>"))', "<b>Salom</b>", "Salom", "&lt;b&gt;Salom&lt;/b&gt"),
    P('from urllib.robotparser import RobotFileParser\nr = RobotFileParser()\nr.parse(["User-agent: *", "Disallow: /"])\nprint(r.can_fetch("bot", "https://x.uz/sahifa"))', "True", "/", "Xato"),
    P('import re\nprint(re.findall(r\'href="([^"]+)"\', \'<a href="/a">A</a> <a href="/b">B</a>\'))', "['/a']", "['<a href=\"/a\">', '<a href=\"/b\">']", "2", lv=2),
    Q("Saytdan maʼlumot yigʻishdan oldin nimani tekshirish kerak?", "robots.txt va saytning foydalanish shartlarini", "Saytning rangini", "Server qaysi tilda yozilganini", "Hech narsani"),
    Q("Sayt rasmiy API taqdim etsa, nima afzal?", "API — u barqaror va ruxsat etilgan yoʻl", "Scraping — u tezroq", "Ikkalasi ham bir xil", "Saytni nusxalab olish"),
    F("Havolalar", "`havolalar(html)` — `html.parser` bilan barcha `<a href>` qiymatlarini (kelish tartibida) qaytarsin.",
      "from html.parser import HTMLParser\n\n\ndef havolalar(html):\n    class H(HTMLParser):\n        def __init__(self):\n            super().__init__()\n            self.r = []\n\n        def handle_starttag(self, tag, attrs):\n            if tag == 'a':\n                href = dict(attrs).get('href')\n                if href:\n                    self.r.append(href)\n    h = H()\n    h.feed(html)\n    return h.r",
      'havolalar(\'<a href="/a">A</a><p><a href="https://x.uz">X</a></p><a>yoʻq</a>\')', 'havolalar("<p>matn</p>")', 'havolalar(\'<A HREF="/katta">K</A>\')', lv=1),
    F("Sarlavha", "`sahifa_nomi(html)` — `<title>` ichidagi matnni (boʻsh joylari tozalangan) qaytarsin; boʻlmasa `None`.",
      "from html.parser import HTMLParser\n\n\ndef sahifa_nomi(html):\n    class T(HTMLParser):\n        def __init__(self):\n            super().__init__()\n            self.ichida = False\n            self.matn = None\n\n        def handle_starttag(self, tag, attrs):\n            self.ichida = tag == 'title'\n\n        def handle_endtag(self, tag):\n            if tag == 'title':\n                self.ichida = False\n\n        def handle_data(self, data):\n            if self.ichida:\n                self.matn = (self.matn or '') + data\n    t = T()\n    t.feed(html)\n    return ' '.join(t.matn.split()) if t.matn else None",
      'sahifa_nomi("<html><head><title>  Mirae   Academy </title></head></html>")', 'sahifa_nomi("<p>yoʻq</p>")', 'sahifa_nomi("<title>A &amp; B</title>")', lv=2),
    F("Toʻliq havolalar", "`toliq(asos, havolalar)` — `urljoin` bilan toʻliq URL larga aylantirib, takrorsiz va `#` bilan boshlanuvchilarsiz (kelish tartibida) qaytarsin.",
      "from urllib.parse import urljoin\n\n\ndef toliq(asos, havolalar):\n    natija = []\n    for h in havolalar:\n        if h.startswith('#'):\n            continue\n        u = urljoin(asos, h)\n        if u not in natija:\n            natija.append(u)\n    return natija",
      'toliq("https://mirae.uz/dars/", ["1", "/kurslar", "#top", "1", "https://python.org"])', 'toliq("https://a.uz", [])', 'toliq("https://a.uz/x/y", ["../z"])'),
    F("Narxlar", "`narxlar(html)` — `class=\"narx\"` boʻlgan teglardagi matndan raqamlarni ajratib (`12 500 soʻm` → `12500`), sonlar roʻyxatini qaytarsin.",
      "import re\nfrom html.parser import HTMLParser\n\n\ndef narxlar(html):\n    class N(HTMLParser):\n        def __init__(self):\n            super().__init__()\n            self.ichida = 0\n            self.r = []\n\n        def handle_starttag(self, tag, attrs):\n            if 'narx' in (dict(attrs).get('class') or '').split():\n                self.ichida += 1\n\n        def handle_endtag(self, tag):\n            if self.ichida:\n                self.ichida -= 1\n\n        def handle_data(self, data):\n            if self.ichida and re.search(r'\\d', data):\n                self.r.append(int(re.sub(r'\\D', '', data)))\n    n = N()\n    n.feed(html)\n    return n.r",
      'narxlar(\'<div><span class="narx">12 500 soʻm</span><span class="nom">Non</span><b class="narx eski">99 000</b></div>\')', 'narxlar("<p>yoʻq</p>")', 'narxlar(\'<i class="narx">5</i>\')', lv=3),
    F("robots.txt", "`ruxsat(robots, yollar, agent=\"MiraeBot\")` — `RobotFileParser` bilan har bir yoʻl uchun ruxsat bor-yoʻqligini `{yoʻl: bool}` qilib qaytarsin (asos: `https://site.uz`).",
      "from urllib.robotparser import RobotFileParser\n\n\ndef ruxsat(robots, yollar, agent='MiraeBot'):\n    r = RobotFileParser()\n    r.parse(robots.splitlines())\n    return {y: r.can_fetch(agent, 'https://site.uz' + y) for y in yollar}",
      'ruxsat("User-agent: *\\nDisallow: /admin/\\n", ["/", "/admin/x", "/blog"])', 'ruxsat("User-agent: MiraeBot\\nDisallow: /\\n", ["/a"])', 'ruxsat("", ["/x"])', lv=2),
    F("Jadvaldan maʼlumot", "`jadval(html)` — birinchi `<table>` dagi `<tr>` qatorlarini `[[katakchalar matni]]` roʻyxati qilib qaytarsin (`th` va `td` ikkalasi ham).",
      "from html.parser import HTMLParser\n\n\ndef jadval(html):\n    class J(HTMLParser):\n        def __init__(self):\n            super().__init__()\n            self.qatorlar = []\n            self.katak = None\n\n        def handle_starttag(self, tag, attrs):\n            if tag == 'tr':\n                self.qatorlar.append([])\n            elif tag in ('td', 'th'):\n                self.katak = ''\n\n        def handle_endtag(self, tag):\n            if tag in ('td', 'th') and self.katak is not None:\n                self.qatorlar[-1].append(self.katak.strip())\n                self.katak = None\n\n        def handle_data(self, data):\n            if self.katak is not None:\n                self.katak += data\n    j = J()\n    j.feed(html)\n    return j.qatorlar",
      'jadval("<table><tr><th>Ism</th><th>Ball</th></tr><tr><td>Ali</td><td>85</td></tr><tr><td>Madina</td><td> 92 </td></tr></table>")', 'jadval("<p></p>")', 'jadval("<table><tr><td>a</td></tr></table>")', lv=3),
    F("Faqat matn", "`faqat_matn(html)` — barcha teglarni olib tashlab, faqat koʻrinadigan matnni qaytarsin: `<script>` va `<style>` ichidagilar tashlansin, boʻsh joylar bittaga qisqartirilsin.",
      "from html.parser import HTMLParser\n\n\ndef faqat_matn(html):\n    class M(HTMLParser):\n        def __init__(self):\n            super().__init__()\n            self.yashirin = 0\n            self.qismlar = []\n\n        def handle_starttag(self, tag, attrs):\n            if tag in ('script', 'style'):\n                self.yashirin += 1\n\n        def handle_endtag(self, tag):\n            if tag in ('script', 'style') and self.yashirin:\n                self.yashirin -= 1\n\n        def handle_data(self, data):\n            if not self.yashirin:\n                self.qismlar.append(data)\n    m = M()\n    m.feed(html)\n    return ' '.join(' '.join(m.qismlar).split())",
      'faqat_matn("<h1>Salom</h1><script>alert(1)</script><p>Mirae   <b>Academy</b></p><style>p{}</style>")', 'faqat_matn("")', 'faqat_matn("<p>5 &lt; 7</p>")', lv=2),
    L("docs", "py-html-parser"),
    L("lc", "html-entity-parser"),
    L("exercism", "markdown", note="Teskari yoʻnalish: Markdown → HTML. Teglar tuzilishini yaxshi tushunasiz."),
    L("docs", "requests-quickstart"),
])

# ---------------------------------------------------------------- 12.6
TOPIC("brauzer-avtomatika", "Brauzer avtomatlashtirish", "Selenium va Playwright: bosish, yozish, kutish", 20, r'''
Baʼzi sahifalar JavaScript bilan yuklanadi yoki faqat tugma bosilgandan keyin maʼlumot koʻrsatadi — `requests` bunda ojiz. **Brauzer avtomatlashtirish** — haqiqiy brauzerni dastur orqali boshqarish: sahifa ochish, matn yozish, tugma bosish, skrinshot olish.

Ikki mashhur vosita: **Selenium** (klassik) va **Playwright** (zamonaviy, tez).

## Selenium
```py norun
from selenium import webdriver               # pip install selenium
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

driver = webdriver.Chrome()
driver.get("https://mirae.uz/kirish")

driver.find_element(By.NAME, "login").send_keys("ali")
driver.find_element(By.NAME, "parol").send_keys("••••••")
driver.find_element(By.CSS_SELECTOR, "button[type=submit]").click()

# Element paydo boʻlguncha (koʻpi bilan 10 soniya) kutish
salom = WebDriverWait(driver, 10).until(EC.visibility_of_element_located((By.ID, "salom")))
print(salom.text)
driver.save_screenshot("kabinet.png")
driver.quit()
```

## Playwright
```py norun
from playwright.sync_api import sync_playwright     # pip install playwright && playwright install

with sync_playwright() as p:
    brauzer = p.chromium.launch(headless=True)
    sahifa = brauzer.new_page()
    sahifa.goto("https://mirae.uz/kurslar")
    sahifa.fill("input[name=q]", "python")
    sahifa.click("text=Qidirish")
    sahifa.wait_for_selector(".kurs-karta")
    print(sahifa.locator(".kurs-karta h3").all_inner_texts())
    sahifa.screenshot(path="kurslar.png", full_page=True)
    brauzer.close()
```

## Lokatorlar
| Usul | Misol | Izoh |
|---|---|---|
| ID | `#salom` | eng ishonchli |
| CSS | `button.asosiy`, `input[name=q]` | universal |
| Matn | `text=Kirish` | foydalanuvchi koʻradigani |
| XPath | `//div[@class='karta']/h3` | kuchli, lekin oʻqish qiyin |

## Kutish (wait) — eng muhim tushuncha
Sahifa elementlari darhol paydo boʻlmaydi. `time.sleep(5)` — yomon (sekin va baribir ishonchsiz). To'g'risi — **shart bajarilguncha kutish**:

```py
def kut(shart, max_soniya=10, qadam=0.5, soat=None):
    """Shart rost boʻlguncha kutadi. Soat — sinash uchun soxta vaqt."""
    t = 0.0
    while t <= max_soniya:
        if shart(t):
            return t
        t += qadam
    raise TimeoutError(f"{max_soniya} soniyada shart bajarilmadi")

# Element 2.3 soniyada paydo boʻladi deb tasavvur qilamiz
print("Topildi:", kut(lambda t: t >= 2.3), "soniyada")
try:
    kut(lambda t: False, max_soniya=3)
except TimeoutError as e:
    print("TimeoutError:", e)
```

!> Avtomatlashtirishni faqat **ruxsat etilgan** joyda ishlating: oʻz saytingizni testlash, ochiq maʼlumotlar, oʻz hisobingizdagi takroriy ishlar. Boshqa birovning tizimiga ruxsatsiz kirish — qonunbuzarlik.

## Esda tuting
- Selenium/Playwright — JavaScript'li sahifalar va formalar uchun.
- `sleep` emas, **shartli kutish**; lokator sifatida ID va CSS.
''', [
    P('print("button[type=submit]".split("[")[0])', "button[type=submit]", "submit", "type"),
    P('print("#salom".startswith("#"), ".karta".startswith("#"))', "True True", "False False", "Xato"),
    P('t = 0\nwhile t < 2.3:\n    t += 0.5\nprint(t)', "2.3", "2.0", "3"),
    P('print(round(sum([0.5] * 5), 1))', "0.5", "5", "2.25"),
    P('import time\nprint(callable(time.sleep))', "False", "None", "sleep"),
    P('x = "//div[@class=\'karta\']/h3"\nprint(x.count("/"))', "1", "2", "4"),
    P('def kut(f, m):\n    t = 0\n    while t <= m:\n        if f(t):\n            return t\n        t += 1\n    raise TimeoutError\nprint(kut(lambda t: t * t > 10, 9))', "3", "10", "⚠ TimeoutError", lv=2),
    Q("Nega `time.sleep(5)` emas, shartli kutish yaxshi?", "Element tayyor boʻlishi bilan davom etadi va sekin internetda ham ishonchli", "sleep Python'da yoʻq", "Shartli kutish kamroq kod", "Farqi yoʻq"),
    Q("Qaysi lokator odatda eng barqaror?", "Element ID si (#salom)", "Elementning ekrandagi koordinatasi", "Uzun XPath", "Sahifadagi tartib raqami"),
    F("Kutish", "`kut(paydo_boladi, max_soniya, qadam=0.5)` — soxta vaqt `t` 0 dan `qadam` bilan oshadi; `t >= paydo_boladi` boʻlganda `t` ni qaytarsin; `t > max_soniya` boʻlsa `TimeoutError`.",
      "def kut(paydo_boladi, max_soniya, qadam=0.5):\n    t = 0.0\n    while t <= max_soniya:\n        if t >= paydo_boladi:\n            return t\n        t += qadam\n    raise TimeoutError()",
      "kut(2.3, 10)", "kut(0, 5)", "kut(7, 3)", "kut(1, 10, 0.25)", lv=1),
    F("CSS selektor", "`selektor(teg, id=None, sinf=None, **atr)` — CSS selektor yasasin: `teg` + `#id` + `.sinf` + `[k=v]` (atributlar alifbo boʻyicha).",
      "def selektor(teg, id=None, sinf=None, **atr):\n    s = teg\n    if id:\n        s += '#' + id\n    if sinf:\n        s += '.' + sinf\n    for k, v in sorted(atr.items()):\n        s += f'[{k}={v}]'\n    return s",
      'selektor("button", sinf="asosiy", type="submit")', 'selektor("div", id="salom")', 'selektor("input", name="q", placeholder="Qidirish")'),
    F("Forma maydonlari", "`maydonlar(html)` — `<input>` va `<textarea>` larning `name` atributlarini (boʻlsa) kelish tartibida qaytarsin (`html.parser`).",
      "from html.parser import HTMLParser\n\n\ndef maydonlar(html):\n    class M(HTMLParser):\n        def __init__(self):\n            super().__init__()\n            self.r = []\n\n        def handle_starttag(self, tag, attrs):\n            if tag in ('input', 'textarea'):\n                n = dict(attrs).get('name')\n                if n:\n                    self.r.append(n)\n\n        handle_startendtag = handle_starttag\n    m = M()\n    m.feed(html)\n    return m.r",
      'maydonlar(\'<form><input name="login"><input type="password" name="parol"/><textarea name="izoh"></textarea><input type="submit"></form>\')', 'maydonlar("<p></p>")', 'maydonlar(\'<input name="a"/>\')', lv=2),
    F("Ssenariy", "`bajar(qadamlar)` — `[(\"och\", url), (\"yoz\", maydon, matn), (\"bos\", tugma)]` qadamlarini \"bajarib\", jurnal satrlari roʻyxatini qaytarsin: `\"OCHILDI url\"`, `\"YOZILDI maydon=matn\"` (parol maydonida matn `***`), `\"BOSILDI tugma\"`; nomaʼlum qadam — `ValueError`.",
      "def bajar(qadamlar):\n    jurnal = []\n    for q in qadamlar:\n        if q[0] == 'och':\n            jurnal.append(f'OCHILDI {q[1]}')\n        elif q[0] == 'yoz':\n            matn = '***' if 'parol' in q[1] else q[2]\n            jurnal.append(f'YOZILDI {q[1]}={matn}')\n        elif q[0] == 'bos':\n            jurnal.append(f'BOSILDI {q[1]}')\n        else:\n            raise ValueError(q[0])\n    return jurnal",
      'bajar([("och", "https://mirae.uz/kirish"), ("yoz", "login", "ali"), ("yoz", "parol", "123"), ("bos", "Kirish")])', "bajar([])", 'bajar([("sakra", 1)])', lv=2),
    F("Skrinshot nomi", "`skrinshot_nomi(url, vaqt)` — URL domeni va yoʻlidan fayl nomi yasasin: `https://mirae.uz/kurslar/python` + `datetime(2026,10,13,9,5)` → `\"mirae.uz_kurslar_python_20261013-0905.png\"` (yoʻl boʻsh boʻlsa faqat domen).",
      "from urllib.parse import urlparse\n\n\ndef skrinshot_nomi(url, vaqt):\n    u = urlparse(url)\n    qismlar = [u.netloc] + [p for p in u.path.split('/') if p]\n    return '_'.join(qismlar) + vaqt.strftime('_%Y%m%d-%H%M') + '.png'",
      "skrinshot_nomi('https://mirae.uz/kurslar/python', __import__('datetime').datetime(2026, 10, 13, 9, 5))",
      "skrinshot_nomi('https://mirae.uz/', __import__('datetime').datetime(2026, 1, 1, 0, 0))",
      "skrinshot_nomi('http://localhost:8000/a', __import__('datetime').datetime(2026, 12, 31, 23, 59))"),
    F("Brauzer tarixi", "`Tarix` sinfi: `och(url)` (oldinga tarixni oʻchiradi), `orqaga()` va `oldinga()` — joriy URL ni qaytaradi (chegaradan chiqmasin). `joriy` atributi.",
      "class Tarix:\n    def __init__(self, bosh):\n        self.r = [bosh]\n        self.i = 0\n\n    @property\n    def joriy(self):\n        return self.r[self.i]\n\n    def och(self, url):\n        self.r = self.r[:self.i + 1] + [url]\n        self.i += 1\n\n    def orqaga(self):\n        self.i = max(0, self.i - 1)\n        return self.joriy\n\n    def oldinga(self):\n        self.i = min(len(self.r) - 1, self.i + 1)\n        return self.joriy",
      '(t := Tarix("a"), t.och("b"), t.och("c"), t.orqaga())[-1]', '(t := Tarix("a"), t.och("b"), t.orqaga(), t.och("x"), t.oldinga())[-1]', 'Tarix("bosh").orqaga()', lv=3),
    F("Qayta urinish", "`ishonchli_bos(natijalar, urinish=3)` — `natijalar` — har bir bosish urinishi natijasi (`True` — muvaffaqiyat). Birinchi muvaffaqiyatli urinish raqamini (1 dan) qaytarsin; `urinish` ta urinishda ham boʻlmasa `TimeoutError`.",
      "def ishonchli_bos(natijalar, urinish=3):\n    for i, ok in enumerate(natijalar[:urinish], 1):\n        if ok:\n            return i\n    raise TimeoutError('element bosilmadi')",
      "ishonchli_bos([False, True, True])", "ishonchli_bos([True])", "ishonchli_bos([False, False, False, True])", "ishonchli_bos([False, False, False, True], urinish=4)", lv=1),
    L("docs", "selenium-first"),
    L("lc", "design-browser-history", note="Brauzerning Orqaga/Oldinga tugmalari mantiqi."),
    L("docs", "py-html-parser"),
    L("lc", "design-a-text-editor"),
])

# ---------------------------------------------------------------- 12.7
TOPIC("pdf", "PDF bilan ishlash", "pypdf: matn ajratish, birlashtirish, sahifalarni tanlash", 20, r'''
Shartnomalar, hisobotlar, cheklar, sertifikatlar — koʻpincha PDF. **pypdf** kutubxonasi bilan PDF ni oʻqish, birlashtirish, boʻlish va aylantirish mumkin.

## Matnni ajratib olish
```py norun
from pypdf import PdfReader                     # pip install pypdf

reader = PdfReader("hisobot.pdf")
print(len(reader.pages), "sahifa")
for i, sahifa in enumerate(reader.pages, 1):
    matn = sahifa.extract_text()
    print(f"--- {i}-sahifa ---")
    print(matn[:200])
```

## Birlashtirish va tanlash
```py norun
from pypdf import PdfReader, PdfWriter

writer = PdfWriter()
for fayl in ["muqova.pdf", "hisobot.pdf", "ilovalar.pdf"]:
    writer.append(fayl)                         # butun faylni qoʻshish
writer.write("toʻliq_hisobot.pdf")

# Faqat 1-3 va 5-sahifalar
r = PdfReader("hisobot.pdf")
w = PdfWriter()
for i in [0, 1, 2, 4]:
    w.add_page(r.pages[i])
w.write("qisqa.pdf")
```

> Yangi PDF **yaratish** uchun (sertifikat, chek) `reportlab` yoki `fpdf2` kutubxonalari ishlatiladi.

## Sahifa oraligʻini tahlil qilish
Printer oynasidagi "1-3, 5, 8-10" kabi yozuvni sahifa roʻyxatiga aylantirish — juda amaliy vazifa:

```py
def sahifalar(matn, jami):
    natija = []
    for qism in matn.replace(" ", "").split(","):
        if "-" in qism:
            a, b = map(int, qism.split("-"))
            natija.extend(range(a, b + 1))
        elif qism:
            natija.append(int(qism))
    return [s for s in natija if 1 <= s <= jami]

print(sahifalar("1-3, 5, 8-10", 9))
```

## PDF dan olingan matnni tozalash
PDF dan chiqqan matn koʻpincha "iflos": qator oxiridagi chiziqchalar, ortiqcha boʻsh joylar, sahifa raqamlari.

```py
import re

xom = """Mirae Academy dasturlash kur-
si talabalari uchun hisobot.

       Sahifa 1 / 3
Natijalar   juda    yaxshi."""

matn = re.sub(r"-\n", "", xom)                         # soʻz boʻlinishini tiklash
matn = re.sub(r"^\s*Sahifa \d+ / \d+\s*$", "", matn, flags=re.M)
matn = re.sub(r"\s+", " ", matn).strip()
print(matn)
```

## Esda tuting
- `PdfReader(...).pages[i].extract_text()` — matn, `PdfWriter` — yangi PDF.
- Sahifa raqamlari foydalanuvchi uchun 1 dan, Python'da indeks 0 dan.
''', [
    P('print(list(range(3, 5 + 1)))', "[3, 4]", "[3, 4, 5, 6]", "[5, 4, 3]"),
    P('print([i - 1 for i in [1, 3, 5]])', "[1, 3, 5]", "[2, 4, 6]", "Xato"),
    P('import re\nprint(re.sub(r"-\\n", "", "kur-\\nsi"))', "kur-si", "kur si", "kur-\\nsi"),
    P('import re\nprint(re.sub(r"\\s+", " ", "a   b\\n\\nc").strip())', "a b\\n\\nc", "abc", "a  b c"),
    P('qism = "8-10"\na, b = map(int, qism.split("-"))\nprint(b - a + 1)', "2", "8-10", "Xato"),
    P('print(sorted(set([3, 1, 3, 2])))', "[3, 1, 2]", "{1, 2, 3}", "[1, 3, 3, 2]"),
    P('sahifalar = ["a", "b", "c", "d"]\nprint(sahifalar[::-1][:2])', "['a', 'b']", "['c', 'd']", "Xato", lv=2),
    Q("pypdf'da birinchi sahifa qanday olinadi?", "reader.pages[0]", "reader.pages[1]", "reader.page(1)", "reader.first()"),
    Q("PDF dan olingan matnda `kur-\\nsi` kabi boʻlinishni tiklashning oddiy usuli?", "re.sub(r\"-\\n\", \"\", matn)", "matn.upper()", "matn.split()", "matn[::-1]"),
    F("Sahifa oraligʻi", "`sahifalar(matn, jami)` — `\"1-3, 5, 8-10\"` → sahifa raqamlari (takrorsiz, saralangan, 1..jami oraligʻida).",
      "def sahifalar(matn, jami):\n    r = set()\n    for q in matn.replace(' ', '').split(','):\n        if '-' in q:\n            a, b = map(int, q.split('-'))\n            r.update(range(a, b + 1))\n        elif q:\n            r.add(int(q))\n    return sorted(s for s in r if 1 <= s <= jami)",
      'sahifalar("1-3, 5, 8-10", 9)', 'sahifalar("2, 2, 1", 5)', 'sahifalar("", 3)', 'sahifalar("4-6", 3)', lv=1),
    F("Matnni tozalash", "`tozala(xom)` — boʻlingan soʻzlarni tiklasin (`-\\n`), `Sahifa N / M` qatorlarini olib tashlasin, barcha boʻsh joylarni bittaga qisqartirsin.",
      "import re\n\n\ndef tozala(xom):\n    m = re.sub(r'-\\n', '', xom)\n    m = re.sub(r'^\\s*Sahifa \\d+ / \\d+\\s*$', '', m, flags=re.M)\n    return re.sub(r'\\s+', ' ', m).strip()",
      'tozala("Mirae kur-\\nsi.\\n  Sahifa 1 / 3\\nNatija   yaxshi.")', 'tozala("")', 'tozala("bir\\nikki")', lv=2),
    F("Qidiruv", "`qidir(sahifalar, soz)` — sahifa matnlari roʻyxatidan `soz` (katta-kichik farqsiz) uchragan sahifa raqamlari (1 dan).",
      "def qidir(sahifalar, soz):\n    return [i for i, m in enumerate(sahifalar, 1) if soz.lower() in m.lower()]",
      'qidir(["Kirish", "Python asoslari", "PYTHON va Django"], "python")', "qidir([], 'x')", 'qidir(["a"], "b")', lv=1),
    F("Boʻlib yuborish", "`bolaklar(jami, hajm)` — `jami` sahifali PDF ni `hajm` sahifadan boʻlaklarga boʻlish rejasi: `[(boshi, oxiri), ...]` (1 dan, oxirgisi qisqa boʻlishi mumkin).",
      "def bolaklar(jami, hajm):\n    return [(b, min(b + hajm - 1, jami)) for b in range(1, jami + 1, hajm)]",
      "bolaklar(10, 4)", "bolaklar(3, 5)", "bolaklar(0, 2)"),
    F("Birlashtirish rejasi", "`reja(fayllar)` — `[(nom, sahifalar_soni)]` dan birlashtirilgan PDF da har bir fayl qaysi sahifalarda turishini `{nom: (boshi, oxiri)}` qilib qaytarsin.",
      "def reja(fayllar):\n    d, b = {}, 1\n    for nom, n in fayllar:\n        d[nom] = (b, b + n - 1)\n        b += n\n    return d",
      'reja([("muqova.pdf", 1), ("hisobot.pdf", 12), ("ilova.pdf", 3)])', "reja([])", 'reja([("a.pdf", 5)])'),
    F("Mundarija", "`mundarija(sahifalar)` — har bir sahifa matnining birinchi boʻsh boʻlmagan qatori sarlavha deb, `\"Sarlavha .... N\"` qatorlarini qaytarsin (sarlavha va raqam orasini nuqtalar bilan 30 belgigacha toʻldiring: `f\"{s:.<28}{n:>2}\"`).",
      "def mundarija(sahifalar):\n    r = []\n    for n, m in enumerate(sahifalar, 1):\n        q = [x.strip() for x in m.splitlines() if x.strip()]\n        if q:\n            r.append(f'{q[0]:.<28}{n:>2}')\n    return r",
      'mundarija(["Kirish\\nmatn", "\\n  Asoslar\\n...", "", "Xulosa"])', "mundarija([])", 'mundarija(["  Yagona bob  "])', lv=2),
    F("Teskari tartib", "`chop_tartibi(n)` — ikki tomonlama chop etish uchun: avval toq sahifalar (1, 3, 5…), keyin juft sahifalar teskari tartibda.",
      "def chop_tartibi(n):\n    return list(range(1, n + 1, 2)) + list(range(n - n % 2, 0, -2))",
      "chop_tartibi(6)", "chop_tartibi(5)", "chop_tartibi(1)", lv=2),
    L("docs", "pypdf-extract"),
    L("docs", "pypdf-merge"),
    L("exercism", "word-count"),
    L("lc", "text-justification", note="Qiyin: matnni sahifa kengligiga tekislash."),
])

# ---------------------------------------------------------------- 12.8
TOPIC("excel", "Excel jadvallari bilan ishlash", "openpyxl: kataklar, formulalar, uslublar", 20, r'''
Buxgalteriya, ombor, baholar jurnali — hamma joyda Excel. **openpyxl** yordamida `.xlsx` fayllarni Python'dan oʻqish va yaratish mumkin.

## Yaratish
```py norun
from openpyxl import Workbook                # pip install openpyxl
from openpyxl.styles import Font

wb = Workbook()
ws = wb.active
ws.title = "Baholar"
ws.append(["Ism", "Matematika", "Fizika", "Oʻrtacha"])
for ism, m, f in [("Ali", 85, 78), ("Madina", 92, 95), ("Bek", 70, 88)]:
    ws.append([ism, m, f])
for r in range(2, ws.max_row + 1):
    ws[f"D{r}"] = f"=AVERAGE(B{r}:C{r})"      # Excel formulasi
ws["A1"].font = Font(bold=True)
ws.column_dimensions["A"].width = 14
wb.save("baholar.xlsx")
```

## Oʻqish
```py norun
from openpyxl import load_workbook

wb = load_workbook("savdo.xlsx", data_only=True)   # formulalar oʻrniga natijalar
ws = wb["Oktabr"]
for qator in ws.iter_rows(min_row=2, values_only=True):
    mahsulot, narx, soni = qator
    print(mahsulot, narx * soni)
```

## Katak manzillari: A1, B12, AA3
Excel ustunlarni harflar bilan belgilaydi: A..Z, keyin AA, AB, ... Bu — 26 lik sanoq tizimiga oʻxshash:

```py
def ustun_raqami(harflar):
    n = 0
    for h in harflar:
        n = n * 26 + (ord(h) - ord("A") + 1)
    return n

def ustun_harfi(n):
    s = ""
    while n:
        n, q = divmod(n - 1, 26)
        s = chr(ord("A") + q) + s
    return s

for x in ["A", "Z", "AA", "AZ", "XFD"]:
    print(x, "→", ustun_raqami(x))
print([ustun_harfi(n) for n in [1, 26, 27, 52, 703]])
```

> `XFD` = 16384 — Excel'dagi oxirgi ustun!

## Excel oʻrniga CSV
Oddiy jadval almashish uchun CSV yetarli va Excel uni toʻgʻridan-toʻgʻri ochadi. Oʻzbekcha harflar Excel'da toʻgʻri chiqishi uchun `encoding="utf-8-sig"` ishlating:

```py
import csv

with open("baholar.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Ism", "Ball"])
    w.writerow(["Oʻktam", 91])
print(open("baholar.csv", "rb").read()[:3])     # BOM — Excel uchun belgi
```

## Esda tuting
- `Workbook()` → `ws.append(qator)` → `wb.save()`; `load_workbook(..., data_only=True)` — oʻqish.
- Ustun harflari ↔ raqamlari: 26 lik tizim (A=1).
''', [
    P('print(ord("C") - ord("A") + 1)', "2", "C", "67"),
    P('n = 0\nfor h in "AB":\n    n = n * 26 + (ord(h) - 64)\nprint(n)', "3", "12", "52"),
    P('print(divmod(27 - 1, 26))', "(1, 1)", "(0, 26)", "1"),
    P('import re\nprint(re.fullmatch(r"([A-Z]+)(\\d+)", "AB12").groups())', "['AB', '12']", "('A', 'B12')", "AB12"),
    P('print(f"=SUM(B2:B{1 + 3})")', "=SUM(B2:B13)", "=SUM(B2:B{4})", "Xato"),
    P('print("utf-8-sig".endswith("sig"), "\\ufeff".encode("utf-8"))', "True b''", "False b'\\ufeff'", "Xato", lv=2),
    P('qator = ("Non", 4000, 3)\nm, n, s = qator\nprint(n * s)', "4000", "Non4000", "Xato"),
    Q("Excel'da `AA` ustuni nechanchi ustun?", "27", "26", "52", "11"),
    Q("`load_workbook(..., data_only=True)` nima beradi?", "Formulalar oʻrniga ularning hisoblangan natijalarini", "Faqat birinchi varaqni", "Faylni faqat oʻqish rejimida ochadi, boshqa farqi yoʻq", "Rasmlarni"),
    F("Ustun raqami", "`ustun_raqami(harflar)` — `\"A\"` → 1, `\"Z\"` → 26, `\"AA\"` → 27.",
      "def ustun_raqami(harflar):\n    n = 0\n    for h in harflar.upper():\n        n = n * 26 + ord(h) - 64\n    return n",
      'ustun_raqami("A")', 'ustun_raqami("AZ")', 'ustun_raqami("XFD")', 'ustun_raqami("zz")', lv=1),
    F("Ustun harfi", "`ustun_harfi(n)` — teskari: 1 → `\"A\"`, 28 → `\"AB\"`, 703 → `\"AAA\"`.",
      "def ustun_harfi(n):\n    s = ''\n    while n:\n        n, q = divmod(n - 1, 26)\n        s = chr(65 + q) + s\n    return s",
      "ustun_harfi(1)", "ustun_harfi(28)", "ustun_harfi(703)", "ustun_harfi(16384)", lv=2),
    F("Katak manzili", "`katak(manzil)` — `\"B12\"` → `(12, 2)` (qator, ustun raqami); notoʻgʻri manzil — `ValueError`.",
      "import re\n\n\ndef katak(manzil):\n    m = re.fullmatch(r'([A-Z]+)([1-9]\\d*)', manzil.upper())\n    if not m:\n        raise ValueError(manzil)\n    n = 0\n    for h in m.group(1):\n        n = n * 26 + ord(h) - 64\n    return (int(m.group(2)), n)",
      'katak("B12")', 'katak("aa1")', 'katak("A0")', 'katak("12B")'),
    F("SUM oraligʻi", "`summa(jadval, oraliq)` — `jadval` — `{\"B2\": 5, ...}` lugʻat; `\"B2:C3\"` oraligʻidagi sonli kataklar yigʻindisi (yoʻq katak — 0).",
      "import re\n\n\ndef summa(jadval, oraliq):\n    def tahlil(m):\n        h, r = re.fullmatch(r'([A-Z]+)(\\d+)', m).groups()\n        n = 0\n        for c in h:\n            n = n * 26 + ord(c) - 64\n        return int(r), n\n    def harf(n):\n        s = ''\n        while n:\n            n, q = divmod(n - 1, 26)\n            s = chr(65 + q) + s\n        return s\n    a, b = oraliq.split(':')\n    (r1, c1), (r2, c2) = tahlil(a), tahlil(b)\n    return sum(jadval.get(f'{harf(c)}{r}', 0) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))",
      'summa({"B2": 5, "C2": 7, "B3": 1, "D9": 100}, "B2:C3")', 'summa({}, "A1:Z99")', 'summa({"A1": 1, "A2": 2, "A3": 3}, "A2:A3")', lv=3),
    F("Excel uchun CSV", "`excel_csv(nom, qatorlar)` — CSV ni `utf-8-sig` bilan yozsin va faylning birinchi 3 baytini (`bytes`) qaytarsin (BOM boʻlishi kerak).",
      "import csv\n\n\ndef excel_csv(nom, qatorlar):\n    with open(nom, 'w', newline='', encoding='utf-8-sig') as f:\n        csv.writer(f).writerows(qatorlar)\n    with open(nom, 'rb') as f:\n        return f.read(3)",
      'excel_csv("a.csv", [["Ism", "Ball"], ["Oʻktam", 91]])', 'excel_csv("b.csv", [])', 'excel_csv("c.csv", [["x"]]) == b"\\xef\\xbb\\xbf"'),
    F("Ustunlar oʻrtachasi", "`ortachalar(qatorlar)` — birinchi qator sarlavha; har bir sonli ustunning oʻrtachasini (1 xona) `{sarlavha: oʻrtacha}` qaytarsin (birinchi ustun — ism, hisobga olinmaydi).",
      "def ortachalar(qatorlar):\n    sar, *data = qatorlar\n    return {s: round(sum(q[i] for q in data) / len(data), 1) for i, s in enumerate(sar) if i > 0} if data else {}",
      'ortachalar([["Ism", "Matematika", "Fizika"], ["Ali", 85, 78], ["Madina", 92, 95], ["Bek", 70, 88]])', 'ortachalar([["Ism", "X"]])', 'ortachalar([["Ism", "A"], ["x", 1]])'),
    F("Ustun kengligi", "`kengliklar(qatorlar)` — har bir ustun uchun eng uzun qiymat (`str` koʻrinishida) uzunligi + 2 — openpyxl'da `column_dimensions[...].width` uchun. Natija: `{\"A\": 8, \"B\": 5, ...}`.",
      "def kengliklar(qatorlar):\n    d = {}\n    for q in qatorlar:\n        for i, v in enumerate(q):\n            h = chr(ord('A') + i)\n            d[h] = max(d.get(h, 0), len(str(v)) + 2)\n    return d",
      'kengliklar([["Ism", "Ball"], ["Madina", 92], ["Ali", 100]])', "kengliklar([])", 'kengliklar([["", None]])', lv=2),
    L("lc", "excel-sheet-column-number"),
    L("lc", "excel-sheet-column-title"),
    L("lc", "design-spreadsheet", note="Kataklar va =X+Y formulalari bilan oddiy elektron jadval."),
    L("docs", "py-csv"),
])

# ---------------------------------------------------------------- 12.9
TOPIC("cqs", "Command Query Separation printsipi", "Holatni oʻzgartiradigan va maʼlumot beradigan metodlarni ajratish", 15, r'''
**CQS** (Command–Query Separation, "buyruq va soʻrovni ajratish") — Bertrand Meyer taklif qilgan oddiy, lekin kuchli qoida:

- **Soʻrov (query)** — maʼlumot **qaytaradi**, hech narsani **oʻzgartirmaydi**;
- **Buyruq (command)** — holatni **oʻzgartiradi**, maʼlumot **qaytarmaydi** (`None`).

> "Savol berish javobni oʻzgartirmasligi kerak."

## Yomon misol
```py
class Navbat:
    def __init__(self):
        self.odamlar = ["Ali", "Vali", "Gani"]

    def keyingi(self):              # ❌ ham soʻrov, ham buyruq
        return self.odamlar.pop(0)

n = Navbat()
if n.keyingi() == "Ali":           # faqat tekshirmoqchi edik...
    print("Ali birinchi")
print(n.odamlar)                    # ...lekin Ali navbatdan chiqib ketdi!
```

## Yaxshi misol
```py
class Navbat:
    def __init__(self):
        self.odamlar = ["Ali", "Vali", "Gani"]

    def birinchi(self):             # soʻrov — oʻzgartirmaydi
        return self.odamlar[0] if self.odamlar else None

    def chiqar(self):               # buyruq — oʻzgartiradi, qaytarmaydi
        self.odamlar.pop(0)

n = Navbat()
if n.birinchi() == "Ali":
    print("Ali birinchi")
print(n.odamlar)                    # hech narsa buzilmadi
n.chiqar()
print(n.birinchi())
```

## Python'da CQS
Python'ning oʻzi ham shu qoidaga koʻp joyda amal qiladi:

```py
royxat = [3, 1, 2]
print(royxat.sort())        # buyruq → None
print(sorted(royxat))       # soʻrov → yangi roʻyxat, asl oʻzgarmaydi
print(royxat.count(2))      # soʻrov
```

| Buyruq (None qaytaradi) | Soʻrov (qiymat qaytaradi) |
|---|---|
| `list.sort()` | `sorted(list)` |
| `list.append(x)` | `len(list)`, `x in list` |
| `dict.update(...)` | `dict.get(k)` |

## Istisno: pop() va boshqalar
`list.pop()`, `dict.setdefault()` — ham oʻzgartiradi, ham qaytaradi. Bular qulaylik uchun ataylab qilingan istisnolar. Oʻz kodingizda esa CQS ni buzish **sababli** boʻlsin.

## Nega muhim?
- Soʻrovlarni istalgancha chaqirsa boʻladi — xavfsiz (testlar, loglar, debugger).
- Kodni oʻqiganda qaysi metod holatni oʻzgartirishi darhol aniq.
- Xatolar kamayadi: "tekshiruv" hech qachon "tasodifiy oʻzgarish" boʻlmaydi.

## Esda tuting
- Soʻrov — `return` bor, holat oʻzgarmaydi; buyruq — holat oʻzgaradi, `return` yoʻq.
''', [
    P('a = [3, 1, 2]\nprint(a.sort())', "[1, 2, 3]", "[3, 1, 2]", "Xato"),
    P('a = [3, 1, 2]\nb = sorted(a)\nprint(a, b)', "[1, 2, 3] [1, 2, 3]", "[3, 1, 2] [3, 1, 2]", "None [1, 2, 3]"),
    P('a = [1, 2, 3]\nif a.pop(0) == 1:\n    pass\nprint(a)', "[1, 2, 3]", "[1, 2]", "Xato"),
    P('d = {"a": 1}\nprint(d.update({"b": 2}), len(d))', "{'a': 1, 'b': 2} 2", "None 1", "2"),
    P('d = {}\nx = d.setdefault("k", [])\nx.append(1)\nprint(d)', "{}", "{'k': []}", "Xato", lv=2),
    P('class H:\n    def __init__(self):\n        self.n = 0\n    def oshir(self):\n        self.n += 1\nh = H()\nprint(h.oshir(), h.n)', "1 1", "0 1", "Xato"),
    P('s = "salom"\nprint(s.upper(), s)', "SALOM SALOM", "salom SALOM", "None salom"),
    Q("CQS ga koʻra soʻrov (query) metod nima qilishi kerak?", "Qiymat qaytarishi, lekin holatni oʻzgartirmasligi", "Holatni oʻzgartirishi va None qaytarishi", "Ham oʻzgartirishi, ham qaytarishi", "Faqat xato chiqarishi"),
    Q("Qaysi biri buyruq (command)?", "list.append(x)", "len(list)", "sorted(list)", "list.count(x)"),
    F("Navbat (CQS)", "`Navbat` sinfi: `qosh(ism)` va `chiqar()` — buyruqlar (`None` qaytaradi); `birinchi()` va `uzunlik()` — soʻrovlar.",
      "class Navbat:\n    def __init__(self):\n        self.r = []\n\n    def qosh(self, ism):\n        self.r.append(ism)\n\n    def chiqar(self):\n        if self.r:\n            self.r.pop(0)\n\n    def birinchi(self):\n        return self.r[0] if self.r else None\n\n    def uzunlik(self):\n        return len(self.r)",
      '(n := Navbat(), n.qosh("Ali"), n.qosh("Vali"), n.birinchi(), n.birinchi(), n.uzunlik())[3:]', '(n := Navbat(), n.qosh("A"), n.chiqar(), n.birinchi())[-1]', '(n := Navbat(), n.qosh("A"))[-1]', lv=1),
    F("Hisob (CQS)", "`Hisob(balans)`: buyruqlar `toldir(s)`, `yech(s)` (yetmasa `ValueError`); soʻrovlar `balans()` va `yecha_oladimi(s)`.",
      "class Hisob:\n    def __init__(self, balans):\n        self._b = balans\n\n    def toldir(self, s):\n        self._b += s\n\n    def yech(self, s):\n        if s > self._b:\n            raise ValueError('mablagʻ yetarli emas')\n        self._b -= s\n\n    def balans(self):\n        return self._b\n\n    def yecha_oladimi(self, s):\n        return s <= self._b",
      "(h := Hisob(100), h.yecha_oladimi(150), h.balans())[1:]", "(h := Hisob(100), h.toldir(50), h.yech(120), h.balans())[-1]", "Hisob(10).yech(20)", "(h := Hisob(5), h.toldir(1))[-1]"),
    F("Ajratib yozing", "`Stek` sinfi CQS bilan: `qosh(x)`, `olib_tashla()` — buyruqlar; `tepa()` — tepadagi element (boʻsh boʻlsa `None`), `boshmi()` — soʻrovlar.",
      "class Stek:\n    def __init__(self):\n        self.r = []\n\n    def qosh(self, x):\n        self.r.append(x)\n\n    def olib_tashla(self):\n        if self.r:\n            self.r.pop()\n\n    def tepa(self):\n        return self.r[-1] if self.r else None\n\n    def boshmi(self):\n        return not self.r",
      "(s := Stek(), s.qosh(1), s.qosh(2), s.tepa(), s.tepa())[3:]", "(s := Stek(), s.qosh(1), s.olib_tashla(), s.boshmi())[-1]", "Stek().tepa()"),
    F("Tekshiruvchi", "`cqs_buzadimi(f, obyekt, *args)` — metodni (nomi bilan) chaqirishdan oldin va keyin `vars(obyekt)` ni solishtiring: holat oʻzgarib, natija `None` boʻlmasa `True` (CQS buzildi).",
      "import copy\n\n\ndef cqs_buzadimi(f, obyekt, *args):\n    oldin = copy.deepcopy(vars(obyekt))\n    natija = getattr(obyekt, f)(*args)\n    return vars(obyekt) != oldin and natija is not None\n\n\nclass Royxat:\n    def __init__(self):\n        self.r = [1, 2, 3]\n\n    def pop(self):\n        return self.r.pop()\n\n    def oxirgi(self):\n        return self.r[-1]\n\n    def qosh(self, x):\n        self.r.append(x)",
      'cqs_buzadimi("pop", Royxat())', 'cqs_buzadimi("oxirgi", Royxat())', 'cqs_buzadimi("qosh", Royxat(), 5)', lv=3),
    F("Savat (CQS)", "`Savat`: buyruqlar `qosh(nom, narx)`, `ochir(nom)`, `tozala()`; soʻrovlar `jami()`, `bormi(nom)`, `nomlar()` (alifbo boʻyicha).",
      "class Savat:\n    def __init__(self):\n        self.d = {}\n\n    def qosh(self, nom, narx):\n        self.d[nom] = narx\n\n    def ochir(self, nom):\n        self.d.pop(nom, None)\n\n    def tozala(self):\n        self.d.clear()\n\n    def jami(self):\n        return sum(self.d.values())\n\n    def bormi(self, nom):\n        return nom in self.d\n\n    def nomlar(self):\n        return sorted(self.d)",
      '(s := Savat(), s.qosh("non", 4000), s.qosh("sut", 12000), s.ochir("non"), (s.jami(), s.nomlar()))[-1]', '(s := Savat(), s.qosh("a", 1), s.tozala(), s.bormi("a"))[-1]', "Savat().jami()"),
    F("Hisoblagich", "`Bosuvchi` sinfi: `bos()` (buyruq), `qiymat()` (soʻrov), `nolla()` (buyruq). `ssenariy(n)` funksiyasi n marta bosib, qiymatni qaytarsin.",
      "class Bosuvchi:\n    def __init__(self):\n        self._n = 0\n\n    def bos(self):\n        self._n += 1\n\n    def qiymat(self):\n        return self._n\n\n    def nolla(self):\n        self._n = 0\n\n\ndef ssenariy(n):\n    b = Bosuvchi()\n    for _ in range(n):\n        b.bos()\n    return b.qiymat()",
      "ssenariy(5)", "ssenariy(0)", "Bosuvchi().bos()"),
    F("Toʻgʻri yoki notoʻgʻri", "`turi(nomlar)` — metod nomlari roʻyxatini CQS boʻyicha ajratsin: `get_`, `is_`, `has_`, `count` bilan boshlanadiganlar → `\"soʻrov\"`, `set_`, `add`, `remove`, `clear`, `update` bilan boshlanadiganlar → `\"buyruq\"`, qolganlari → `\"nomaʼlum\"`. Natija: `{nom: tur}`.",
      "def turi(nomlar):\n    d = {}\n    for n in nomlar:\n        if n.startswith(('get_', 'is_', 'has_', 'count')):\n            d[n] = 'soʻrov'\n        elif n.startswith(('set_', 'add', 'remove', 'clear', 'update')):\n            d[n] = 'buyruq'\n        else:\n            d[n] = 'nomaʼlum'\n    return d",
      'turi(["get_name", "add_item", "is_empty", "process"])', "turi([])", 'turi(["clear", "count_items", "update"])', lv=1),
    L("lc", "min-stack", note="Soʻrovlar (top, getMin) va buyruqlar (push, pop) aniq ajratilgan."),
    L("lc", "design-circular-deque"),
    L("exercism", "bank-account"),
    L("lc", "smallest-number-in-infinite-set"),
])

# ---------------------------------------------------------------- 12.10
TOPIC("numpy", "NumPy asoslari", "Massivlar, vektor amallar va tez hisob-kitob", 25, r'''
**NumPy** — Python'da ilmiy hisoblashlar, maʼlumotlar tahlili va sunʼiy intellektning poydevori. Uning asosiy obyekti — **ndarray**: bir xil turdagi sonlar massivi, ustida amallar C tilidagi tezlikda bajariladi.

```py norun
import numpy as np                       # pip install numpy

a = np.array([1, 2, 3, 4])
print(a * 10)                            # [10 20 30 40]  — siklsiz!
print(a.mean(), a.sum(), a.max())        # 2.5 10 4

m = np.array([[1, 2, 3], [4, 5, 6]])
print(m.shape)                           # (2, 3)
print(m.T)                               # transpozitsiya
print(m.sum(axis=0))                     # ustunlar yigʻindisi: [5 7 9]
print(m[m > 3])                          # maska: [4 5 6]
print(np.arange(0, 10, 2))               # [0 2 4 6 8]
print(np.zeros((2, 3)))
```

> Bu ilova ichida NumPy yoʻq (internet talab qilinadi), shuning uchun uning gʻoyalarini **sof Python** bilan qayta yaratamiz — keyin NumPy'ni oʻrnatganingizda nima "sehr" boʻlayotganini aniq bilasiz.

## Vektorlashtirish gʻoyasi
NumPy'da `a * 10` — massivning **har bir** elementiga amal. Sof Python'da buni roʻyxat ifodasi bilan qilamiz:

```py
a = [1, 2, 3, 4]
b = [10, 20, 30, 40]
print([x * 10 for x in a])               # a * 10
print([x + y for x, y in zip(a, b)])     # a + b
print(sum(x * y for x, y in zip(a, b)))  # np.dot(a, b)
print(sum(a) / len(a))                   # a.mean()
```

## Shakl (shape) va reshape
```py
def shakl(m):
    return (len(m), len(m[0])) if m and isinstance(m[0], list) else (len(m),)

def reshape(royxat, q, u):
    assert q * u == len(royxat), "elementlar soni mos emas"
    return [royxat[i * u:(i + 1) * u] for i in range(q)]

m = reshape(list(range(6)), 2, 3)
print(m, shakl(m))
print([list(t) for t in zip(*m)])        # m.T
print([sum(ustun) for ustun in zip(*m)]) # m.sum(axis=0)
print([sum(q) for q in m])               # m.sum(axis=1)
```

## Maska (mantiqiy indekslash)
```py
ballar = [85, 42, 91, 67, 58, 99]
maska = [b >= 60 for b in ballar]        # ballar >= 60
print(maska)
print([b for b, m in zip(ballar, maska) if m])   # ballar[ballar >= 60]
print(sum(maska), "kishi oʻtdi")         # (ballar >= 60).sum()
```

## Nega NumPy tez?
| Sof Python roʻyxati | NumPy massivi |
|---|---|
| har element — alohida Python obyekti | bitta uzluksiz xotira bloki |
| sikl Python'da | sikl C'da |
| 1 mln element yigʻindisi ~50 ms | ~1 ms |

## Esda tuting
- NumPy: `np.array`, amallar butun massivga (`a * 2`, `a + b`), `shape`, `axis`, maska.
- Sof Python ekvivalenti: roʻyxat ifodalari, `zip`, `sum`.
''', [
    P('a = [1, 2, 3]\nprint([x * 2 for x in a])', "[1, 2, 3, 1, 2, 3]", "[1, 2, 3, 2]", "6"),
    P('a = [1, 2, 3]\nprint(a * 2)', "[2, 4, 6]", "[1, 2, 3]", "Xato", lv=2,
      why="Sof Python'da `roʻyxat * 2` — roʻyxatni takrorlash! NumPy'da esa har elementni 2 ga koʻpaytirish. Bu farq koʻp chalkashlikka sabab boʻladi."),
    P('a, b = [1, 2], [3, 4]\nprint(sum(x * y for x, y in zip(a, b)))', "[3, 8]", "10", "21"),
    P('m = [[1, 2, 3], [4, 5, 6]]\nprint([sum(u) for u in zip(*m)])', "[6, 15]", "21", "[[1, 4], [2, 5], [3, 6]]"),
    P('print(list(range(0, 10, 3)))', "[0, 3, 6, 9, 10]", "[3, 6, 9]", "[0, 3, 6]"),
    P('b = [5, 80, 61]\nprint(sum(x >= 60 for x in b))', "141", "[80, 61]", "True"),
    P('r = list(range(6))\nprint([r[i * 3:(i + 1) * 3] for i in range(2)])', "[[0, 1], [2, 3], [4, 5]]", "[0, 1, 2, 3, 4, 5]", "[[0, 3], [1, 4], [2, 5]]"),
    Q("NumPy'da `a * 10` nima qiladi?", "Massivning har bir elementini 10 ga koʻpaytiradi", "Massivni 10 marta takrorlaydi", "Xato beradi", "Faqat birinchi elementni koʻpaytiradi"),
    Q("`m.sum(axis=0)` (2×3 matritsa) nima qaytaradi?", "Har bir ustun yigʻindisi (3 ta son)", "Har bir qator yigʻindisi (2 ta son)", "Barcha elementlar yigʻindisi", "Matritsa shaklini"),
    F("Vektor amallar", "`qosh(a, b)` — elementma-element yigʻindi; uzunliklar har xil boʻlsa `ValueError` (NumPy kabi).",
      "def qosh(a, b):\n    if len(a) != len(b):\n        raise ValueError('shakllar mos emas')\n    return [x + y for x, y in zip(a, b)]",
      "qosh([1, 2, 3], [10, 20, 30])", "qosh([], [])", "qosh([1], [1, 2])", lv=1),
    F("Skalyar koʻpaytma", "`dot(a, b)` — `np.dot` ekvivalenti (1D vektorlar).",
      "def dot(a, b):\n    if len(a) != len(b):\n        raise ValueError('shakllar mos emas')\n    return sum(x * y for x, y in zip(a, b))",
      "dot([1, 2, 3], [4, 5, 6])", "dot([], [])", "dot([2], [3, 4])", lv=1),
    F("reshape", "`reshape(royxat, q, u)` — 1D roʻyxatni `q×u` matritsaga; elementlar soni mos kelmasa `ValueError`. `u=-1` boʻlsa ustunlar soni avtomatik hisoblansin.",
      "def reshape(royxat, q, u):\n    if u == -1:\n        if len(royxat) % q:\n            raise ValueError('mos emas')\n        u = len(royxat) // q\n    if q * u != len(royxat):\n        raise ValueError('mos emas')\n    return [royxat[i * u:(i + 1) * u] for i in range(q)]",
      "reshape(list(range(6)), 2, 3)", "reshape(list(range(6)), 3, -1)", "reshape([1, 2, 3], 2, 2)", lv=2),
    F("axis boʻyicha", "`yigindi(m, axis=None)` — `None` → barcha elementlar, `0` → ustunlar, `1` → qatorlar.",
      "def yigindi(m, axis=None):\n    if axis is None:\n        return sum(sum(q) for q in m)\n    if axis == 0:\n        return [sum(u) for u in zip(*m)]\n    if axis == 1:\n        return [sum(q) for q in m]\n    raise ValueError(axis)",
      "yigindi([[1, 2, 3], [4, 5, 6]])", "yigindi([[1, 2, 3], [4, 5, 6]], 0)", "yigindi([[1, 2, 3], [4, 5, 6]], 1)", "yigindi([[1]], 2)", lv=2),
    F("Maska", "`maska_filtr(a, chegara)` — `(maska, tanlanganlar, soni)` kortejini qaytarsin: maska — `x > chegara` roʻyxati.",
      "def maska_filtr(a, chegara):\n    m = [x > chegara for x in a]\n    t = [x for x, k in zip(a, m) if k]\n    return (m, t, len(t))",
      "maska_filtr([85, 42, 91, 67], 60)", "maska_filtr([], 0)", "maska_filtr([1, 2], 5)"),
    F("Statistika", "`tavsif(a)` — `{\"min\", \"max\", \"mean\", \"std\"}` (`std` — populyatsiya standart chetlanishi, `np.std` kabi; mean va std 3 xona).",
      "def tavsif(a):\n    n = len(a)\n    o = sum(a) / n\n    s = (sum((x - o) ** 2 for x in a) / n) ** 0.5\n    return {'min': min(a), 'max': max(a), 'mean': round(o, 3), 'std': round(s, 3)}",
      "tavsif([2, 4, 4, 4, 5, 5, 7, 9])", "tavsif([5])", "tavsif([1, 2, 3])", lv=2),
    F("Normallashtirish", "`normallash(a)` — min-max normallashtirish: `(x - min) / (max - min)` (3 xona); hamma qiymat teng boʻlsa — nollar.",
      "def normallash(a):\n    lo, hi = min(a), max(a)\n    if hi == lo:\n        return [0.0] * len(a)\n    return [round((x - lo) / (hi - lo), 3) for x in a]",
      "normallash([10, 20, 30, 40])", "normallash([5, 5])", "normallash([3, 1, 2])", lv=2),
    L("docs", "numpy-beginners"),
    L("docs", "numpy-quickstart"),
    L("lc", "reshape-the-matrix"),
    L("lc", "transpose-matrix"),
])
